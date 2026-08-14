from io import BytesIO

from openpyxl import load_workbook
import pandas as pd

from core import schedule_core
from app.services.export_excel_service import ExcelExportRequest, build_export_excel_bytes


class FakeExportCore:
    def __init__(self):
        self.kwargs = None

    def build_empty_schedule_excel_bytes(self, **kwargs):
        self.kwargs = kwargs
        return b"xlsx-bytes"


def test_build_export_excel_bytes_delegates_to_schedule_core():
    core = FakeExportCore()
    calendar_df = pd.DataFrame([{"date": "2026-03-16", "week": 3, "weekday": "周一"}])
    timetable_df = pd.DataFrame([{"period": 1, "start": "08:10", "end": "08:55"}])
    blocks_df = pd.DataFrame([{"name": "张三", "week": 3}])
    all_slot_df = pd.DataFrame([{"week": 3, "free_members": "李四"}])

    data = build_export_excel_bytes(
        schedule_core=core,
        request=ExcelExportRequest(
            occupancy={(3, "周一", 1): {"张三"}},
            students=["张三", "李四"],
            weeks=[3],
            calendar_df=calendar_df,
            timetable_df=timetable_df,
            blocks_df=blocks_df,
            all_slot_df=all_slot_df,
            threshold=0,
        ),
    )

    assert data == b"xlsx-bytes"
    assert core.kwargs["students"] == ["张三", "李四"]
    assert core.kwargs["weeks"] == [3]
    assert core.kwargs["calendar_df"] is calendar_df
    assert core.kwargs["timetable_df"] is timetable_df
    assert core.kwargs["blocks_df"] is blocks_df
    assert core.kwargs["all_slot_df"] is all_slot_df
    assert core.kwargs["threshold"] == 0


def test_schedule_core_export_builds_weekly_matrix_workbook():
    data = schedule_core.build_empty_schedule_excel_bytes(
        occupancy={
            (6, "周一", 1): {"张三"},
            (6, "周一", 2): {"李四"},
            (6, "周一", 12): {"王五"},
            (6, "周二", 1): {"王五"},
            (6, "周五", 9): {"张三"},
        },
        students=["张三", "李四", "王五"],
        weeks=[6],
        calendar_df=pd.DataFrame(),
        timetable_df=schedule_core.default_timetable(),
        blocks_df=pd.DataFrame(),
        all_slot_df=None,
    )

    workbook = load_workbook(BytesIO(data))
    assert workbook.sheetnames[:5] == ["第6周", "空课明细", "成员完整性检查", "课程占用明细", "处理日志摘要"]

    sheet = workbook["第6周"]
    assert sheet["A1"].value == "第六周空课表"
    assert "A1:U2" in {str(item) for item in sheet.merged_cells.ranges}
    assert "B3:E4" in {str(item) for item in sheet.merged_cells.ranges}
    assert "B5:E18" in {str(item) for item in sheet.merged_cells.ranges}
    assert sheet["B3"].value == "周一"
    assert sheet["A5"].value == "1-2节"
    assert sheet["A19"].value == "3-4节"
    assert sheet["A33"].value == "中午"
    assert sheet["A47"].value == "5-6节"
    assert sheet["A75"].value == "9-11节"
    assert sheet["B5"].value == "王五"
    assert sheet["F5"].value == "张三，李四"
    assert sheet["B19"].value == "张三，李四，王五"
    assert sheet["B33"].value == "张三，李四"
    assert sheet["R75"].value == "李四，王五"


def test_schedule_core_export_only_shows_member_names():
    data = schedule_core.build_empty_schedule_excel_bytes(
        occupancy={(6, "周一", 1): {"刘金富｜办公室｜部长"}},
        students=["刘金富｜办公室｜部长", "高怡昕｜办公室｜部长"],
        weeks=[6],
        calendar_df=pd.DataFrame(),
        timetable_df=schedule_core.default_timetable(),
        blocks_df=pd.DataFrame(
            [{"name": "刘金富", "member_key": "刘金富｜办公室｜部长", "course": "课程A"}]
        ),
        all_slot_df=pd.DataFrame(
            [{
                "week": 6,
                "weekday": "周一",
                "period": 1,
                "time": "08:10-08:55",
                "free_count": 1,
                "free_members": "高怡昕｜办公室｜部长",
                "occupied_count": 1,
                "occupied_members": "刘金富｜办公室｜部长",
            }]
        ),
        file_df=pd.DataFrame([{"成员": "刘金富｜办公室｜部长"}]),
    )

    workbook = load_workbook(BytesIO(data))
    values = [
        str(cell.value)
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    ]

    assert any("刘金富" in value for value in values)
    assert any("高怡昕" in value for value in values)
    assert all("办公室" not in value and "部长" not in value for value in values)


def test_schedule_core_export_neutralizes_formula_like_imported_text():
    data = schedule_core.build_empty_schedule_excel_bytes(
        occupancy={(6, "周一", 1): {"=HYPERLINK(\"https://example.invalid\")"}},
        students=["=HYPERLINK(\"https://example.invalid\")", "+SUM(1,1)"],
        weeks=[6],
        calendar_df=pd.DataFrame(),
        timetable_df=schedule_core.default_timetable(),
        blocks_df=pd.DataFrame(
            [{"name": "@user", "course": "-10", "message": "\t=NOW()"}]
        ),
        all_slot_df=pd.DataFrame(
            [{"free_members": "=HYPERLINK(\"https://example.invalid\")"}]
        ),
    )

    workbook = load_workbook(BytesIO(data), data_only=False)
    values = [
        cell.value
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str)
    ]

    assert values
    assert all(cell.data_type != "f" for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row)
    assert "'=HYPERLINK(\"https://example.invalid\")" in values
    assert "'+SUM(1,1)" in values
