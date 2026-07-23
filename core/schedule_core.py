from __future__ import annotations

import io
import contextlib
import hashlib
import json
import logging
import os
import re
import unicodedata
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.worksheet import Worksheet


WEEKDAYS = ["周一", "周二", "周三", "周四", "周五", "周六", "周日"]
WEEKDAY_INDEX = {weekday: index for index, weekday in enumerate(WEEKDAYS)}
SEMESTER_START = date(2026, 3, 2)
TEACHING_WEEKS = 18
PERIOD_DISPLAY_ORDER = [1, 2, 3, 4, 12, 13, 5, 6, 7, 8, 9, 10, 11]
PERIOD_DISPLAY_INDEX = {period: index for index, period in enumerate(PERIOD_DISPLAY_ORDER)}
ENGLISH_WEEKDAY_MAP = {
    "Monday": "周一",
    "Tuesday": "周二",
    "Wednesday": "周三",
    "Thursday": "周四",
    "Friday": "周五",
    "Saturday": "周六",
    "Sunday": "周日",
}

PERIOD_GROUPS = {
    "1-2": [1, 2],
    "3-4": [3, 4],
    "5-6": [5, 6],
    "7-8": [7, 8],
    "9-11": [9, 10, 11],
}

EXPORT_WEEKDAYS = ["周一", "周二", "周三", "周四", "周五"]
EXPORT_PERIOD_BLOCKS = [
    ("1-2节", [1, 2], 5),
    ("3-4节", [3, 4], 19),
    ("中午", [12, 13], 33),
    ("5-6节", [5, 6], 47),
    ("7-8节", [7, 8], 61),
    ("9-11节", [9, 10, 11], 75),
]

NOON_PERIOD_TIMES = {
    12: ("12:40", "13:25"),
    13: ("13:30", "14:15"),
}

NON_CLASS_KEYWORDS = {
    "",
    "-",
    "午",
    "不排课",
    "报到注册",
    "清明节",
    "运动会",
    "端午节",
    "劳动节",
    "校外",
    "教研",
    "期末周",
    "考试周",
}

_OCR_ENGINE = None
LOGGER = logging.getLogger("konggu")
PARSE_CACHE_VERSION = 32
OCR_LAYOUT_CACHE_VERSION = 2
NAME_STOPWORDS = {
    "办公室",
    "外联部",
    "宣传部",
    "活动部",
    "干事",
    "部长",
    "中方课表",
    "英方课表",
    "课表",
    "路人甲",
}


class OCRConfigurationError(RuntimeError):
    """Raised when the local OCR runtime cannot start or execute."""


def _ocr_error_message(exc: BaseException) -> str:
    message = str(exc).strip() or exc.__class__.__name__
    return f"OCR 配置异常：{message}"


def _source_name(source: dict[str, Any]) -> str:
    return str(
        source.get("name")
        or source.get("file_name")
        or Path(str(source.get("path") or source.get("source_path") or "")).name
        or "unknown.pdf"
    )


def _source_path(source: dict[str, Any]) -> str:
    return str(source.get("path") or source.get("source_path") or "")


def _source_bytes(source: dict[str, Any]) -> bytes:
    data = source.get("content")
    if data is None:
        data = source.get("bytes")
    if data is None:
        data = source.get("bytes_data")
    return data if isinstance(data, bytes) else b""


def _source_kind(source: dict[str, Any]) -> str:
    file_name = _source_name(source)
    path = _source_path(source)
    detected = source.get("kind") or infer_pdf_kind(file_name, path)
    if detected:
        return str(detected)
    return infer_pdf_kind("chinese")


def _source_digest(source: dict[str, Any]) -> str:
    data = _source_bytes(source)
    if data:
        return hashlib.sha256(data).hexdigest()
    path = _source_path(source)
    if path and Path(path).exists():
        hasher = hashlib.sha256()
        with Path(path).open("rb") as handle:
            for chunk in iter(lambda: handle.read(1024 * 1024), b""):
                hasher.update(chunk)
        return hasher.hexdigest()
    return hashlib.sha256(_source_name(source).encode("utf-8", errors="ignore")).hexdigest()


def default_timetable() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {"period": 1, "start": "08:10", "end": "08:55"},
            {"period": 2, "start": "09:00", "end": "09:45"},
            {"period": 3, "start": "10:15", "end": "11:00"},
            {"period": 4, "start": "11:05", "end": "11:50"},
            {"period": 12, "start": "12:40", "end": "13:25"},
            {"period": 13, "start": "13:30", "end": "14:15"},
            {"period": 5, "start": "14:30", "end": "15:15"},
            {"period": 6, "start": "15:20", "end": "16:05"},
            {"period": 7, "start": "16:25", "end": "17:10"},
            {"period": 8, "start": "17:15", "end": "18:00"},
            {"period": 9, "start": "19:10", "end": "19:55"},
            {"period": 10, "start": "20:00", "end": "20:45"},
            {"period": 11, "start": "20:50", "end": "21:35"},
        ]
    )


def validate_timetable(timetable: pd.DataFrame) -> tuple[pd.DataFrame, list[str]]:
    required = {"period", "start", "end"}
    if timetable is None or timetable.empty or not required.issubset(timetable.columns):
        return default_timetable(), ["节次时间表缺失，已使用默认时间表。"]
    clean = timetable.copy()
    clean["period"] = clean["period"].astype(int)
    clean["_display_order"] = clean["period"].map(lambda period: PERIOD_DISPLAY_INDEX.get(int(period), 10_000 + int(period)))
    return clean.sort_values("_display_order").drop(columns="_display_order").reset_index(drop=True), []


def period_time_range(period: int) -> str:
    if int(period) in NOON_PERIOD_TIMES:
        start, end = NOON_PERIOD_TIMES[int(period)]
        return f"{start}-{end}"
    timetable = default_timetable().set_index("period")
    row = timetable.loc[int(period)]
    return f"{row['start']}-{row['end']}"


def infer_pdf_kind(file_name: str, path: str = "") -> str:
    haystack = f"{file_name} {path}"
    if "英方" in haystack or "uk" in haystack.lower() or "english" in haystack.lower():
        return "英方"
    if "中方" in haystack or "chinese" in haystack.lower():
        return "中方"
    return ""


def local_pdf_sources_from_dir(root: str | Path) -> list[dict[str, Any]]:
    sources: list[dict[str, Any]] = []
    for path in sorted(Path(root).rglob("*.pdf")):
        kind = infer_pdf_kind(path.name, str(path)) or "中方"
        sources.append(
            {
                "name": path.name,
                "path": str(path),
                "kind": kind,
                "content": b"",
            }
        )
    return sources


def _parse_cache_root() -> Path:
    configured = os.environ.get("KONGGU_PARSE_CACHE")
    if configured:
        return Path(configured)
    return Path(__file__).resolve().parents[1] / "cache" / "parsed_blocks"


def _parse_cache_path(source: dict[str, Any], kind: str) -> Path:
    digest = _source_digest(source)
    name_digest = hashlib.sha1(_source_name(source).encode("utf-8", errors="ignore")).hexdigest()[:12]
    safe_kind = re.sub(r"[^0-9A-Za-z_\-\u4e00-\u9fff]+", "_", kind).strip("_") or "unknown"
    return _parse_cache_root() / f"{digest}_{name_digest}_{safe_kind}.json"


def _load_parse_cache(source: dict[str, Any], kind: str) -> list[dict[str, Any]] | None:
    cache_path = _parse_cache_path(source, kind)
    if not cache_path.exists():
        return None
    try:
        payload = json.loads(cache_path.read_text(encoding="utf-8"))
    except Exception as exc:
        LOGGER.warning("解析缓存读取失败: %s error=%s", cache_path.name, exc)
        return None
    if (
        payload.get("version") != PARSE_CACHE_VERSION
        or payload.get("hash") != _source_digest(source)
        or payload.get("file_name") != _source_name(source)
        or payload.get("kind") != kind
        or not isinstance(payload.get("blocks"), list)
    ):
        return None
    return payload["blocks"]


def _write_parse_cache(source: dict[str, Any], kind: str, blocks: list[dict[str, Any]]) -> None:
    if not blocks:
        return
    cache_path = _parse_cache_path(source, kind)
    payload = {
        "version": PARSE_CACHE_VERSION,
        "hash": _source_digest(source),
        "kind": kind,
        "file_name": _source_name(source),
        "block_count": len(blocks),
        "blocks": blocks,
    }
    try:
        cache_path.parent.mkdir(parents=True, exist_ok=True)
        cache_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    except Exception as exc:
        LOGGER.warning("解析缓存写入失败: %s error=%s", cache_path.name, exc)


def _open_pdf_document(source: dict[str, Any]) -> Any | None:
    try:
        import fitz
    except Exception:
        return None

    path = _source_path(source)
    content = _source_bytes(source)
    try:
        if path and Path(path).exists():
            return fitz.open(path)
        if content:
            return fitz.open(stream=content, filetype="pdf")
    except Exception:
        return None
    return None


def _page_words(page: Any) -> list[dict[str, Any]]:
    return [
        {
            "x0": float(word[0]),
            "y0": float(word[1]),
            "x1": float(word[2]),
            "y1": float(word[3]),
            "text": str(word[4]).replace("\xa0", " ").strip(),
        }
        for word in page.get_text("words")
        if str(word[4]).replace("\xa0", " ").strip()
    ]


def _item_center(item: dict[str, Any]) -> tuple[float, float]:
    return ((float(item["x0"]) + float(item["x1"])) / 2, (float(item["y0"]) + float(item["y1"])) / 2)


def _cluster_items_by_y(items: list[dict[str, Any]], tolerance: float) -> list[list[dict[str, Any]]]:
    groups: list[list[dict[str, Any]]] = []
    for item in sorted(items, key=lambda value: (_item_center(value)[1], value["x0"])):
        center_y = _item_center(item)[1]
        for group in groups:
            group_y = sum(_item_center(value)[1] for value in group) / len(group)
            if abs(group_y - center_y) <= tolerance:
                group.append(item)
                break
        else:
            groups.append([item])
    return groups


def _x_bands(columns: list[dict[str, Any]]) -> list[dict[str, Any]]:
    ordered = sorted(columns, key=lambda column: column["center"])
    if not ordered:
        return []
    centers = [float(column["center"]) for column in ordered]
    gaps = [b - a for a, b in zip(centers, centers[1:]) if b > a]
    edge_pad = max(18.0, (min(gaps) if gaps else 45.0) * 0.55)
    bands: list[dict[str, Any]] = []
    for index, column in enumerate(ordered):
        left = (centers[index - 1] + centers[index]) / 2 if index else centers[index] - edge_pad
        right = (centers[index] + centers[index + 1]) / 2 if index + 1 < len(centers) else centers[index] + edge_pad
        bands.append({**column, "left": left, "right": right})
    return bands


def _pick_x_band(bands: list[dict[str, Any]], x_center: float) -> dict[str, Any] | None:
    for band in bands:
        if band["left"] <= x_center < band["right"]:
            return band
    if not bands:
        return None
    nearest = min(bands, key=lambda band: abs(band["center"] - x_center))
    gap = min((b["right"] - b["left"] for b in bands), default=36.0)
    if abs(nearest["center"] - x_center) <= max(18.0, gap * 0.55):
        return nearest
    return None


def parse_actual_pdf_sources(
    sources: list[dict[str, Any]],
    uploaded_calendar_df: pd.DataFrame | None = None,
) -> tuple[list[dict[str, Any]], pd.DataFrame, list[str], pd.DataFrame]:
    blocks: list[dict[str, Any]] = []
    errors: list[str] = []

    for source in sources:
        file_name = _source_name(source)
        kind = _source_kind(source)
        image_only_pdf = False
        used_ocr = False
        try:
            cached = _load_parse_cache(source, kind)
            if cached is not None:
                LOGGER.info("命中解析缓存: %s blocks=%s", file_name, len(cached))
                blocks.extend(cached)
                continue

            LOGGER.info("解析文件: %s kind=%s", file_name, kind)
            parsed: list[dict[str, Any]] = []
            if kind == "中方":
                parsed = _parse_chinese_pdf_layout(source, file_name)
                if parsed:
                    LOGGER.info("坐标表格解析完成: %s blocks=%s", file_name, len(parsed))
                    parsed = _finalize_parsed_blocks(parsed, file_name, kind)
                    _write_parse_cache(source, kind, parsed)
                    blocks.extend(parsed)
                    continue
                parsed = _parse_chinese_legacy_web_layout(source, file_name)
                if parsed:
                    LOGGER.info("旧版网页课表解析完成: %s blocks=%s", file_name, len(parsed))
                    parsed = _finalize_parsed_blocks(parsed, file_name, kind)
                    _write_parse_cache(source, kind, parsed)
                    blocks.extend(parsed)
                    continue
            else:
                parsed = _parse_english_pdf_grid_layout(source, file_name)
                if parsed:
                    if len(parsed) < 8:
                        try:
                            ocr_parsed = _parse_english_ocr_week_grid(source, file_name)
                        except OCRConfigurationError as exc:
                            LOGGER.warning("英方稀疏坐标解析 OCR 补扫失败: %s error=%s", file_name, exc)
                            ocr_parsed = []
                        if len(ocr_parsed) > len(parsed):
                            LOGGER.info("英方 OCR 补扫替换稀疏结果: %s blocks=%s -> %s", file_name, len(parsed), len(ocr_parsed))
                            for item in ocr_parsed:
                                item.setdefault("text_source", "OCR")
                            parsed = ocr_parsed
                    LOGGER.info("英方坐标表格解析完成: %s blocks=%s", file_name, len(parsed))
                    parsed = _finalize_parsed_blocks(parsed, file_name, kind)
                    _write_parse_cache(source, kind, parsed)
                    blocks.extend(parsed)
                    continue

            text = _extract_pdf_text(source)
            image_only_pdf = _is_image_only_pdf_text(text)
            if image_only_pdf:
                LOGGER.info("检测到图片型 PDF，进入 OCR: %s", file_name)
            if not _is_usable_extracted_text(text, kind):
                LOGGER.info("内嵌文本质量不足，进入 OCR: %s", file_name)
                text = _extract_pdf_ocr_text(source)
                used_ocr = True
            if kind == "中方" and used_ocr:
                embedded_name = _extract_chinese_name_from_text(text)
                filename_name = _extract_name_from_filename(file_name)
                if embedded_name and filename_name and embedded_name != filename_name:
                    errors.append(
                        f"{file_name}：课表正文姓名“{embedded_name}”与文件名姓名“{filename_name}”不一致，"
                        "已跳过自动占用计算。请确认是否传错或复制错课表。"
                    )
                    LOGGER.warning(
                        "中方 OCR 姓名不一致，已跳过: %s embedded=%s filename=%s",
                        file_name,
                        embedded_name,
                        filename_name,
                    )
                    continue
            parsed = _parse_chinese_text(text, file_name) if kind == "中方" else _parse_english_text(text, file_name)
            if not parsed and not used_ocr:
                LOGGER.info("内嵌文本解析为空，回退 OCR: %s", file_name)
                ocr_text = _extract_pdf_ocr_text(source)
                used_ocr = True
                if ocr_text.strip():
                    parsed = _parse_chinese_text(ocr_text, file_name) if kind == "中方" else _parse_english_text(ocr_text, file_name)
            if not parsed:
                parsed = (
                    _parse_chinese_ocr_table_layout(source, file_name)
                    if kind == "中方"
                    else _parse_english_ocr_week_grid(source, file_name)
                )
            elif kind == "中方" and used_ocr and image_only_pdf:
                profile = detect_schedule_layout_profile(source, kind)
                if profile == "cdut_undergrad_full_term_cn":
                    table_parsed = _parse_chinese_ocr_table_layout(source, file_name)
                    parsed = _choose_chinese_ocr_parse_candidate(parsed, table_parsed, file_name)
            if kind != "中方" and parsed and len(parsed) < 8:
                try:
                    ocr_parsed = _parse_english_ocr_week_grid(source, file_name)
                except OCRConfigurationError as exc:
                    LOGGER.warning("英方稀疏文本解析 OCR 补扫失败: %s error=%s", file_name, exc)
                    ocr_parsed = []
                if len(ocr_parsed) > len(parsed):
                    LOGGER.info("英方 OCR 补扫替换稀疏文本结果: %s blocks=%s -> %s", file_name, len(parsed), len(ocr_parsed))
                    for item in ocr_parsed:
                        item.setdefault("text_source", "OCR")
                    parsed = ocr_parsed
                    used_ocr = True
            if not parsed and kind == "中方":
                parsed = _parse_chinese_ocr_legacy_layout(source, file_name)
            if parsed:
                for item in parsed:
                    item.setdefault("text_source", "OCR" if used_ocr else "内嵌文本")
                parsed = _finalize_parsed_blocks(parsed, file_name, kind)
                full_term_issue = ""
                full_term_profile = kind == "中方" and used_ocr and detect_schedule_layout_profile(source, kind) == "cdut_undergrad_full_term_cn"
                if full_term_profile:
                    full_term_issue = _chinese_full_term_ocr_confidence_issue(parsed, used_ocr=used_ocr)
                if full_term_profile:
                    visual_parsed = _parse_chinese_full_term_visual_occupancy(source, file_name)
                    visual_parsed = _finalize_parsed_blocks(visual_parsed, file_name, kind)
                    visual_issue = _chinese_full_term_ocr_confidence_issue(visual_parsed, used_ocr=bool(visual_parsed))
                    if visual_parsed and not visual_issue and (
                        full_term_issue or _should_prefer_chinese_full_term_visual_result(parsed, visual_parsed)
                    ):
                        for item in visual_parsed:
                            item.setdefault("text_source", "视觉表格")
                        LOGGER.info(
                            "中方整学期视觉表格解析替换 OCR: %s issue=%s blocks=%s -> %s",
                            file_name,
                            full_term_issue or "视觉覆盖更完整",
                            len(parsed),
                            len(visual_parsed),
                        )
                        _write_parse_cache(source, kind, visual_parsed)
                        blocks.extend(visual_parsed)
                        continue
                if full_term_issue:
                    errors.append(
                        f"{file_name}：图片型中方整学期课表 OCR 结果低置信（{full_term_issue}），"
                        "已跳过自动占用计算。请提供可复制文字的课表，或在人工校对表中确认。"
                    )
                    LOGGER.warning("低置信中方整学期 OCR 结果已跳过: %s issue=%s blocks=%s", file_name, full_term_issue, len(parsed))
                    continue
                if _is_low_confidence_chinese_ocr_result(parsed, used_ocr=used_ocr):
                    metrics = _course_fragment_metrics(parsed, source="中方")
                    errors.append(
                        f"{file_name}：图片型中方课表 OCR 结果低置信"
                        f"（候选 {metrics['count']} 条，碎片 {metrics['fragment_count']} 条，"
                        f"碎片率 {float(metrics['fragment_ratio']):.0%}），已跳过自动占用计算。"
                        "请提供可复制文字的课表，或在人工校对表中确认。"
                    )
                    LOGGER.warning("低置信中方 OCR 结果已跳过: %s blocks=%s", file_name, len(parsed))
                    continue
                if _is_low_confidence_english_ocr_result(parsed, used_ocr=used_ocr):
                    errors.append(
                        f"{file_name}：图片型英方课表 OCR 覆盖不足，仅识别到 {len(parsed)} 条占用，已跳过自动占用计算。"
                        "请提供可复制文字的课表，或在人工校对表中确认。"
                    )
                    LOGGER.warning("低覆盖英方 OCR 结果已跳过: %s blocks=%s", file_name, len(parsed))
                    continue
                LOGGER.info("文件解析完成: %s blocks=%s", file_name, len(parsed))
                _write_parse_cache(source, kind, parsed)
                blocks.extend(parsed)
            else:
                if image_only_pdf and used_ocr:
                    errors.append(
                        f"{file_name}：图片型 PDF 已启用 OCR，但未识别到可解析的课程占用。"
                        "请重新导出可复制文字的原始 PDF，或检查扫描清晰度。"
                    )
                else:
                    errors.append(f"{file_name}：未识别到课程占用。")
        except OCRConfigurationError as exc:
            context = "图片型 PDF 需要 OCR，但 " if image_only_pdf else "课表需要 OCR，但 "
            errors.append(f"{file_name}：{context}{exc}")
        except Exception as exc:
            errors.append(f"{file_name}：{exc}")

    calendar_df = uploaded_calendar_df if uploaded_calendar_df is not None else synthesize_calendar_from_blocks(blocks)
    preview_df = blocks_to_dataframe(blocks).head(20)
    return blocks, calendar_df, errors, preview_df


def synthesize_calendar_from_blocks(blocks: list[dict[str, Any]]) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    semester_start = _semester_start_date()
    for week in range(1, _teaching_weeks() + 1):
        for index, weekday in enumerate(WEEKDAYS):
            current = semester_start + timedelta(days=(week - 1) * 7 + index)
            rows.append({"date": current.isoformat(), "week": week, "weekday": weekday})
    return pd.DataFrame(rows)


def build_occupancy(blocks: list[dict[str, Any]]) -> dict[tuple[int, str, int], set[str]]:
    occupancy: dict[tuple[int, str, int], set[str]] = defaultdict(set)
    for block in blocks:
        week = block.get("week")
        weekday = block.get("weekday")
        period = block.get("period")
        name = block.get("name")
        if week is None or not weekday or period is None or not name:
            continue
        occupancy[(int(week), str(weekday), int(period))].add(str(name))
    return dict(occupancy)


def build_slot_table(
    occupancy: dict[tuple[int, str, int], set[str]],
    students: list[str],
    weeks: list[int],
    weekdays: list[str],
    periods: list[int],
) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    calendar = synthesize_calendar_from_blocks(
        [{"week": week, "weekday": weekday, "date": _date_for_weekday(week, weekday).isoformat()} for week in weeks for weekday in weekdays]
    )
    date_lookup = {(int(row.week), row.weekday): row.date for row in calendar.itertuples(index=False)}
    all_students = set(students)
    for week in weeks:
        for weekday in weekdays:
            for period in periods:
                occupied = set(occupancy.get((week, weekday, period), set()))
                free = sorted(all_students - occupied)
                rows.append(
                    {
                        "week": week,
                        "date": date_lookup.get((week, weekday), _date_for_weekday(week, weekday).isoformat()),
                        "weekday": weekday,
                        "period": period,
                        "time": period_time_range(period),
                        "free_count": len(free),
                        "free_members": "、".join(free),
                        "occupied_count": len(occupied),
                        "occupied_members": "、".join(sorted(occupied)),
                    }
                )
    return pd.DataFrame(rows)


def blocks_to_dataframe(blocks: list[dict[str, Any]]) -> pd.DataFrame:
    if not blocks:
        return pd.DataFrame(columns=["name", "source", "week", "date", "weekday", "period", "time", "course"])
    return pd.DataFrame(blocks).sort_values(["name", "week", "weekday", "period"]).reset_index(drop=True)


def _dedupe_blocks(blocks: list[dict[str, Any]]) -> list[dict[str, Any]]:
    by_slot: dict[tuple[str, int, str, int], dict[str, Any]] = {}
    for block in blocks:
        key = (
            str(block.get("name", "")),
            int(block.get("week") or 0),
            str(block.get("weekday", "")),
            int(block.get("period") or 0),
        )
        existing = by_slot.get(key)
        if existing is None or _course_quality(block.get("course")) > _course_quality(existing.get("course")):
            by_slot[key] = block
    return list(by_slot.values())


def _finalize_parsed_blocks(blocks: list[dict[str, Any]], file_name: str, kind: str) -> list[dict[str, Any]]:
    cleaned: list[dict[str, Any]] = []
    for item in blocks:
        course = str(item.get("course") or "").strip()
        if _is_empty_or_non_class(course):
            continue

        item = dict(item)
        item["course"] = course
        try:
            day = _parse_date(str(item.get("date") or ""))
        except Exception:
            day = None
        if day is not None:
            item["date"] = day.isoformat()
            item["week"] = _week_from_date(day)
            item["weekday"] = WEEKDAYS[day.weekday()]
        cleaned.append(item)

    finalized = _dedupe_blocks(cleaned)
    for item in finalized:
        item.setdefault("source_file", file_name)
        item.setdefault("kind", kind)
    return finalized


def _course_quality(course: Any) -> int:
    text = str(course or "").strip()
    if not text or text == "课程":
        return 0
    score = len(text)
    if re.search(r"[\u4e00-\u9fff]", text):
        score += 20
    if re.search(r"^[A-Z]\d|^E\d|^C\d|^校外|^\d", text):
        score -= 30
    return score


def _is_low_confidence_chinese_ocr_result(blocks: list[dict[str, Any]], *, used_ocr: bool) -> bool:
    if not used_ocr or not blocks:
        return False
    metrics = _course_fragment_metrics(blocks, source="中方")
    count = int(metrics["count"])
    ratio = float(metrics["fragment_ratio"])
    if count < 12:
        return False
    return ratio >= 0.50 if count < 20 else ratio >= 0.35


def _chinese_full_term_ocr_confidence_issue(blocks: list[dict[str, Any]], *, used_ocr: bool) -> str:
    if not used_ocr:
        return ""
    chinese_blocks = [block for block in blocks if block.get("source") == "中方"]
    if len(chinese_blocks) < 20:
        return f"候选占用仅 {len(chinese_blocks)} 条"

    weeks = sorted({int(block.get("week") or 0) for block in chinese_blocks if int(block.get("week") or 0) > 0})
    if not weeks:
        return "未识别到有效周次"
    if len(chinese_blocks) >= 40 and len(weeks) < 10:
        return f"只覆盖 {len(weeks)} 个离散周次：{','.join(str(week) for week in weeks)}"
    if len(chinese_blocks) >= 80 and len(weeks) < 12:
        return f"高占用结果只覆盖 {len(weeks)} 个周次：{','.join(str(week) for week in weeks)}"
    return ""


def _should_prefer_chinese_full_term_visual_result(
    parsed_blocks: list[dict[str, Any]],
    visual_blocks: list[dict[str, Any]],
) -> bool:
    if not parsed_blocks or not visual_blocks:
        return False
    parsed_weeks = {int(block.get("week") or 0) for block in parsed_blocks if int(block.get("week") or 0) > 0}
    visual_weeks = {int(block.get("week") or 0) for block in visual_blocks if int(block.get("week") or 0) > 0}
    if len(visual_weeks) < max(10, min(len(parsed_weeks), 18)):
        return False
    return len(visual_blocks) >= max(len(parsed_blocks) + 40, int(len(parsed_blocks) * 1.6))


def _choose_chinese_ocr_parse_candidate(
    text_parsed: list[dict[str, Any]],
    table_parsed: list[dict[str, Any]],
    file_name: str,
) -> list[dict[str, Any]]:
    if not text_parsed or not table_parsed:
        return text_parsed or table_parsed

    finalized_text = _finalize_parsed_blocks(text_parsed, file_name, "中方")
    if not _is_low_confidence_chinese_ocr_result(finalized_text, used_ocr=True):
        return text_parsed

    finalized_table = _finalize_parsed_blocks(table_parsed, file_name, "中方")
    if not _is_low_confidence_chinese_ocr_result(finalized_table, used_ocr=True):
        return table_parsed
    return text_parsed


def _is_low_confidence_english_ocr_result(blocks: list[dict[str, Any]], *, used_ocr: bool) -> bool:
    if not used_ocr:
        return False
    english_blocks = [block for block in blocks if block.get("source") == "英方"]
    if not english_blocks:
        return False
    return len(english_blocks) < 8


def _looks_like_low_confidence_course_fragment(course: str) -> bool:
    text = re.sub(r"\s+", "", str(course or "").strip())
    if len(text) <= 1:
        return True
    if _looks_like_plausible_chinese_encoded_course_cell(text):
        return False
    if "*" in text:
        return True
    if re.search(r"\d{3,}-\d|\d{4}-\d{2}|\d{3,}", text):
        return True
    if re.fullmatch(r"[A-Za-z]{0,3}\d{1,4}[A-Za-z0-9-]*", text):
        return True
    if re.fullmatch(r"[A-Za-z]+", text) and len(text) <= 5:
        return True
    return False


def _looks_like_plausible_chinese_encoded_course_cell(text: str) -> bool:
    compact = re.sub(r"\s+", "", str(text or ""))
    if not re.search(r"[\u4e00-\u9fff]", compact):
        return False
    if _is_empty_or_non_class(compact) or _looks_like_room(compact):
        return False
    if re.search(r"(报到注册|运动会|清明节|劳动节|端午节|考试周|期末周|不排|教研|教妍|微研|教好)", compact):
        return False
    if re.search(r"[\u4e00-\u9fff][A-Za-z]\d{1,3}", compact):
        return True
    if re.search(r"[A-Za-z]\d{1,3}[\u4e00-\u9fff]", compact):
        return True
    if re.search(r"[\u4e00-\u9fff]\d{1,2}(?!\d)", compact):
        return True
    return False


def _course_fragment_metrics(blocks: list[dict[str, Any]], *, source: str = "") -> dict[str, Any]:
    scoped = [block for block in blocks if not source or block.get("source") == source]
    fragments = [
        str(block.get("course") or "")
        for block in scoped
        if _looks_like_low_confidence_course_fragment(str(block.get("course") or ""))
    ]
    return {
        "count": len(scoped),
        "fragment_count": len(fragments),
        "fragment_ratio": len(fragments) / max(1, len(scoped)),
        "fragment_samples": fragments[:8],
    }


def build_empty_schedule_excel_bytes(
    occupancy: dict[tuple[int, str, int], set[str]],
    students: list[str],
    weeks: list[int],
    calendar_df: pd.DataFrame,
    timetable_df: pd.DataFrame,
    threshold: int = 0,
    blocks_df: pd.DataFrame | None = None,
    all_slot_df: pd.DataFrame | None = None,
) -> bytes:
    timetable, _ = validate_timetable(timetable_df)
    if all_slot_df is None or all_slot_df.empty:
        all_slot_df = build_slot_table(occupancy, students, weeks, WEEKDAYS, list(timetable["period"]))

    export_df = all_slot_df.copy()
    if threshold:
        export_df = export_df[export_df["free_count"] >= threshold]
    selected_weeks = sorted({int(week) for week in weeks})
    if selected_weeks and "week" in export_df.columns:
        export_df = export_df[export_df["week"].astype(int).isin(selected_weeks)]
    export_df = export_df.rename(
        columns={
            "week": "周次",
            "date": "日期",
            "weekday": "星期",
            "period": "节次",
            "time": "时间",
            "free_count": "空闲人数",
            "free_members": "空闲人员",
            "occupied_count": "有课人数",
            "occupied_members": "有课人员",
        }
    )

    member_rows = []
    for name in students:
        member_rows.append({"成员姓名": name, "状态": "已参与统计"})
    member_df = pd.DataFrame(member_rows)
    logs_df = pd.DataFrame(
        [
            {"项目": "成员数量", "值": len(students)},
            {"项目": "时间格数量", "值": len(all_slot_df)},
            {"项目": "周次工作表数量", "值": len(weeks)},
            {"项目": "生成时间", "值": datetime.now().strftime("%Y-%m-%d %H:%M:%S")},
        ]
    )

    output = io.BytesIO()
    workbook = Workbook()
    ordered_weeks = selected_weeks or _ordered_weeks_from_calendar(calendar_df)
    if ordered_weeks:
        for index, week in enumerate(ordered_weeks):
            sheet = workbook.active if index == 0 else workbook.create_sheet()
            sheet.title = f"第{week}周"
            _populate_weekly_availability_sheet(
                sheet=sheet,
                week=week,
                students=students,
                occupancy=occupancy,
                threshold=threshold,
            )
    else:
        sheet = workbook.active
        sheet.title = "空课表"
        _populate_empty_weekly_sheet(sheet)

    _append_dataframe_sheet(workbook, "空课明细", export_df)
    _append_dataframe_sheet(workbook, "成员完整性检查", member_df)
    if blocks_df is None:
        blocks_df = pd.DataFrame()
    _append_dataframe_sheet(workbook, "课程占用明细", blocks_df)
    _append_dataframe_sheet(workbook, "处理日志摘要", logs_df)

    workbook.save(output)
    return output.getvalue()


def _populate_weekly_availability_sheet(
    *,
    sheet: Worksheet,
    week: int,
    students: list[str],
    occupancy: dict[tuple[int, str, int], set[str]],
    threshold: int,
) -> None:
    """Create the association-style weekly free-time matrix."""

    _prepare_weekly_sheet_grid(sheet)
    sheet["A1"] = f"第{_to_chinese_number(week)}周空课表"

    for day_index, weekday in enumerate(EXPORT_WEEKDAYS):
        start_column = 2 + day_index * 4
        end_column = start_column + 3
        sheet.merge_cells(start_row=3, start_column=start_column, end_row=4, end_column=end_column)
        cell = sheet.cell(row=3, column=start_column)
        cell.value = weekday
        cell.font = Font(name="宋体", size=12)

    for label, periods, start_row in EXPORT_PERIOD_BLOCKS:
        end_row = start_row + 13
        sheet.merge_cells(start_row=start_row, start_column=1, end_row=end_row, end_column=1)
        period_cell = sheet.cell(row=start_row, column=1)
        period_cell.value = label
        period_cell.font = Font(name="宋体", size=11)

        for day_index, weekday in enumerate(EXPORT_WEEKDAYS):
            start_column = 2 + day_index * 4
            end_column = start_column + 3
            sheet.merge_cells(start_row=start_row, start_column=start_column, end_row=end_row, end_column=end_column)
            free_members = _free_members_for_period_block(
                occupancy=occupancy,
                students=students,
                week=week,
                weekday=weekday,
                periods=periods,
            )
            value = "，".join(free_members) if not threshold or len(free_members) >= threshold else ""
            cell = sheet.cell(row=start_row, column=start_column)
            cell.value = value
            cell.font = Font(name="宋体", size=11)


def _populate_empty_weekly_sheet(sheet: Worksheet) -> None:
    _prepare_weekly_sheet_grid(sheet)
    sheet["A1"] = "空课表"


def _prepare_weekly_sheet_grid(sheet: Worksheet) -> None:
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    sheet.merge_cells("A1:U2")
    title = sheet["A1"]
    title.font = Font(name="宋体", size=18)
    title.fill = PatternFill(fill_type="solid", fgColor="FFFFFF")

    sheet.merge_cells("A3:A4")
    for column in range(1, 22):
        sheet.column_dimensions[get_column_letter(column)].width = 13
    max_row = 4 + len(EXPORT_PERIOD_BLOCKS) * 14
    for row in range(1, max_row + 1):
        sheet.row_dimensions[row].height = 18
        for column in range(1, 22):
            cell = sheet.cell(row=row, column=column)
            cell.alignment = alignment
            cell.border = border
    sheet.row_dimensions[1].height = 24
    sheet.row_dimensions[2].height = 24


def _free_members_for_period_block(
    *,
    occupancy: dict[tuple[int, str, int], set[str]],
    students: list[str],
    week: int,
    weekday: str,
    periods: list[int],
) -> list[str]:
    occupied: set[str] = set()
    for period in periods:
        occupied.update(occupancy.get((week, weekday, period), set()))
    return [student for student in students if student not in occupied]


def _append_dataframe_sheet(workbook: Workbook, title: str, dataframe: pd.DataFrame) -> None:
    sheet = workbook.create_sheet(title=title)
    for row in dataframe_to_rows(dataframe, index=False, header=True):
        sheet.append(row)
    sheet.freeze_panes = "A2"
    header_font = Font(name="宋体", size=11, bold=True)
    body_font = Font(name="宋体", size=11)
    alignment = Alignment(vertical="top", wrap_text=True)
    for row in sheet.iter_rows():
        for cell in row:
            cell.font = header_font if cell.row == 1 else body_font
            cell.alignment = alignment
    for column_cells in sheet.columns:
        values = [str(cell.value or "") for cell in column_cells]
        width = min(max(len(value) for value in values) + 2, 42)
        sheet.column_dimensions[column_cells[0].column_letter].width = max(width, 10)


def _ordered_weeks_from_calendar(calendar_df: pd.DataFrame) -> list[int]:
    if calendar_df is None or calendar_df.empty or "week" not in calendar_df.columns:
        return []
    weeks: set[int] = set()
    for value in calendar_df["week"]:
        try:
            weeks.add(int(value))
        except Exception:
            continue
    return sorted(weeks)


def _to_chinese_number(value: int) -> str:
    digits = "零一二三四五六七八九"
    if value <= 0:
        return str(value)
    if value < 10:
        return digits[value]
    if value == 10:
        return "十"
    if value < 20:
        return "十" + digits[value % 10]
    if value < 100:
        tens, ones = divmod(value, 10)
        return digits[tens] + "十" + (digits[ones] if ones else "")
    return str(value)


def _extract_pdf_text(source: dict[str, Any]) -> str:
    doc = _open_pdf_document(source)
    if doc is None:
        return ""
    try:
        return "\n".join(page.get_text("text") for page in doc)
    finally:
        doc.close()


def _is_usable_extracted_text(text: str, kind: str) -> bool:
    """判断 PDF 内嵌文本是否足以直接解析，不够可靠时回退 OCR。"""
    clean = text.strip()
    if len(clean) < 80:
        return False
    chinese_chars = re.findall(r"[\u4e00-\u9fff]", clean)
    if kind == "中方":
        return bool(
            len(chinese_chars) >= 20
            and re.search(r"(?:^|\n)\s*\d{1,2}\s*周", clean)
            and any(keyword in clean for keyword in ("周一", "星期一", "1-2", "3-4", "5-6", "7-8"))
        )
    return bool(
        re.search(r"\d{4}-\d{2}-\d{2}", clean)
        and any(day in clean for day in ENGLISH_WEEKDAY_MAP)
        and re.search(r"(?:^|\n)\s*(?:[1-9]|10|11)\s*(?:\n|$)", clean)
    )


def _is_image_only_pdf_text(text: str) -> bool:
    return len(text.strip()) == 0


def _extract_pdf_ocr_text(source: dict[str, Any]) -> str:
    cache_path = _ocr_cache_path(source)
    if cache_path.exists():
        LOGGER.info("命中 OCR 缓存: %s", _source_name(source))
        return cache_path.read_text(encoding="utf-8", errors="ignore")

    items = _extract_pdf_ocr_items(source)
    pages: dict[int, list[dict[str, Any]]] = defaultdict(list)
    for item in items:
        pages[int(item.get("page", 0))].append(item)
    text = "\n\n".join(
        "\n".join(item["text"] for item in sorted(page_items, key=lambda value: (value["y0"], value["x0"])))
        for _, page_items in sorted(pages.items())
    )
    try:
        cache_path.parent.mkdir(parents=True, exist_ok=True)
        cache_path.write_text(text, encoding="utf-8")
    except Exception as exc:
        LOGGER.warning("OCR 文本缓存写入失败: %s error=%s", cache_path.name, exc)
    return text


def _extract_pdf_ocr_items(source: dict[str, Any]) -> list[dict[str, Any]]:
    cache_path = _ocr_layout_cache_path(source)
    if cache_path.exists():
        try:
            payload = json.loads(cache_path.read_text(encoding="utf-8"))
            if payload.get("version") == OCR_LAYOUT_CACHE_VERSION and isinstance(payload.get("items"), list):
                LOGGER.info("命中 OCR 坐标缓存: %s", _source_name(source))
                return payload["items"]
        except Exception as exc:
            LOGGER.warning("OCR 坐标缓存读取失败: %s error=%s", cache_path.name, exc)

    try:
        import numpy as np
    except Exception:
        return []

    doc = _open_pdf_document(source)
    if doc is None:
        return []

    ocr = _build_ocr_engine()
    items: list[dict[str, Any]] = []
    try:
        for page_number, page in enumerate(doc):
            pix = page.get_pixmap(dpi=120, alpha=False)
            image = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)
            if pix.n > 3:
                image = image[:, :, :3]
            try:
                with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
                    result = ocr.ocr(image) if hasattr(ocr, "ocr") else ocr.predict(image)
            except TypeError:
                with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
                    result = ocr.predict(image)
            except Exception as exc:
                raise OCRConfigurationError(_ocr_error_message(exc)) from exc
            items.extend(_extract_ocr_items_from_result(result, page_number, pix.width, pix.height))
    finally:
        doc.close()

    payload = {
        "version": OCR_LAYOUT_CACHE_VERSION,
        "hash": _source_digest(source),
        "file_name": _source_name(source),
        "items": sorted(items, key=lambda value: (value["page"], value["y0"], value["x0"])),
    }
    try:
        cache_path.parent.mkdir(parents=True, exist_ok=True)
        cache_path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    except Exception as exc:
        LOGGER.warning("OCR 坐标缓存写入失败: %s error=%s", cache_path.name, exc)
    return payload["items"]


def _build_ocr_engine():
    global _OCR_ENGINE
    if _OCR_ENGINE is not None:
        return _OCR_ENGINE

    try:
        from paddleocr import PaddleOCR
    except Exception as exc:
        raise OCRConfigurationError(_ocr_error_message(exc)) from exc

    det_dir = _find_ocr_model_dir("PP-OCRv4_mobile_det")
    rec_dir = _find_ocr_model_dir("PP-OCRv4_mobile_rec")
    kwargs: dict[str, Any] = {
        "use_doc_orientation_classify": False,
        "use_doc_unwarping": False,
        "use_textline_orientation": False,
        "enable_mkldnn": False,
        "enable_hpi": False,
        "device": "cpu",
        "ocr_version": "PP-OCRv4",
        "lang": "ch",
    }
    if det_dir and rec_dir:
        kwargs.update(
            {
                "text_detection_model_name": "PP-OCRv4_mobile_det",
                "text_detection_model_dir": str(det_dir),
                "text_recognition_model_name": "PP-OCRv4_mobile_rec",
                "text_recognition_model_dir": str(rec_dir),
                "lang": None,
                "ocr_version": None,
            }
        )
    elif os.environ.get("KONGGU_ALLOW_OCR_MODEL_DOWNLOAD") != "1":
        raise OCRConfigurationError(
            "OCR 配置异常：扫描件 OCR 模型未就绪。请通过离线安装包修复材料，或将 "
            "PP-OCRv4_mobile_det 和 PP-OCRv4_mobile_rec 放入 %LOCALAPPDATA%\\Konggu\\ocr_models。"
        )
    try:
        with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
            _OCR_ENGINE = PaddleOCR(**kwargs)
    except Exception as exc:
        raise OCRConfigurationError(_ocr_error_message(exc)) from exc
    return _OCR_ENGINE


def check_ocr_runtime(*, run_probe: bool = True) -> dict[str, Any]:
    """Verify that the configured OCR runtime can start and execute locally."""

    try:
        ocr = _build_ocr_engine()
    except Exception as exc:
        return {"ready": False, "stage": "init", "error": str(exc)}

    if run_probe:
        try:
            import numpy as np

            image = np.full((64, 64, 3), 255, dtype=np.uint8)
            with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
                try:
                    ocr.ocr(image) if hasattr(ocr, "ocr") else ocr.predict(image)
                except TypeError:
                    ocr.predict(image)
        except Exception as exc:
            return {"ready": False, "stage": "probe", "error": _ocr_error_message(exc)}

    return {"ready": True, "stage": "ready", "error": ""}


def detect_schedule_layout_profile(source: dict[str, Any], kind: str | None = None) -> str:
    detected_kind = kind or _source_kind(source)
    text = _extract_pdf_text(source)
    items: list[dict[str, Any]] = []
    if len(text.strip()) < 50:
        try:
            items = _extract_pdf_ocr_items(source)
        except Exception:
            items = []
    return _detect_schedule_layout_profile_from_text_items(detected_kind, text, items)


def _detect_schedule_layout_profile_from_text_items(
    kind: str,
    text: str,
    items: list[dict[str, Any]],
) -> str:
    normalized_text = unicodedata.normalize("NFKC", text)
    tokens = [_normalize_ocr_text(str(item.get("text") or "")) for item in items]
    tokens = [unicodedata.normalize("NFKC", token) for token in tokens]
    joined = "\n".join([normalized_text, *tokens])
    compact = re.sub(r"\s+", "", joined)
    if kind == "中方":
        if (
            "成都理工大学本科学生课表" in compact
            and ("周/节" in compact or "周节" in compact or "图节" in compact or "频节" in compact)
            and ("星期" in compact or "星斯" in compact)
            and (
                re.search(r"(?:第)?1\s*[周期期]", compact)
                or len(re.findall(r"\d{2}/\d{2}", compact)) >= 3
                or _looks_like_chinese_slot_header_token(compact)
            )
        ):
            return "cdut_undergrad_full_term_cn"
        if "学期理论课表" in compact or "旧版课表打印" in compact:
            return "cdut_legacy_chinese_web"
    if kind == "英方":
        if (
            ("Timetablefor" in compact or "Timetable for" in joined)
            and ("CDUTSino-British" in compact or "Sino-British" in joined)
            and any(day in joined for day in ENGLISH_WEEKDAY_MAP)
        ):
            return "cdut_sino_british_english_web"
        if "第" in compact and "周" in compact and any(day in compact for day in WEEKDAYS):
            return "cdut_sino_british_english_week_grid"
    return "unknown"


def _ocr_cache_path(source: dict[str, Any]) -> Path:
    digest = _source_digest(source)
    cache_root = Path(os.environ.get("KONGGU_OCR_TEXT_CACHE", "")) if os.environ.get("KONGGU_OCR_TEXT_CACHE") else Path(__file__).resolve().parents[1] / "cache" / "pdf_text"
    return cache_root / f"{digest}.txt"


def _ocr_layout_cache_path(source: dict[str, Any]) -> Path:
    digest = _source_digest(source)
    cache_root = (
        Path(os.environ.get("KONGGU_OCR_LAYOUT_CACHE", ""))
        if os.environ.get("KONGGU_OCR_LAYOUT_CACHE")
        else Path(__file__).resolve().parents[1] / "cache" / "pdf_layout"
    )
    return cache_root / f"{digest}.json"


def _find_ocr_model_dir(model_name: str) -> Path | None:
    roots = [
        Path(os.environ.get("KONGGU_OCR_MODEL_DIR", "")),
        Path(os.environ.get("PADDLE_PDX_CACHE_HOME", "")) / "official_models",
        Path(__file__).resolve().parents[1] / "resources" / "ocr_models",
        Path(__file__).resolve().parents[1] / "models" / "paddleocr",
        Path(__file__).resolve().parents[1] / "cache" / "paddlex" / "official_models",
    ]
    for root in roots:
        if not root or not str(root):
            continue
        candidates = [
            root / model_name,
            root / f"{model_name}_infer",
            root / f"{model_name}_inference",
        ]
        for candidate in candidates:
            if _is_valid_paddle_model_dir(candidate):
                return candidate
    return None


def _is_valid_paddle_model_dir(path: Path) -> bool:
    if not path.exists() or not path.is_dir():
        return False
    names = {item.name for item in path.iterdir()}
    return bool(
        names.intersection({"inference.yml", "inference.yaml", "model.yml", "config.yml"})
        or any(name.endswith(".pdmodel") or name.endswith(".json") for name in names)
    )


def _extract_ocr_lines(result: Any) -> list[str]:
    lines: list[str] = []
    if result is None:
        return lines
    if isinstance(result, dict):
        for key in ("rec_texts", "texts"):
            values = result.get(key)
            if isinstance(values, list):
                lines.extend(str(value) for value in values if value)
        return lines
    if isinstance(result, (list, tuple)):
        for item in result:
            if isinstance(item, dict):
                lines.extend(_extract_ocr_lines(item))
            elif isinstance(item, (list, tuple)):
                if len(item) >= 2 and isinstance(item[1], (list, tuple)) and item[1]:
                    lines.append(str(item[1][0]))
                else:
                    lines.extend(_extract_ocr_lines(item))
            elif item:
                text = str(item)
                if text.strip():
                    lines.append(text)
    return lines


def _extract_ocr_items_from_result(result: Any, page_number: int, width: int, height: int) -> list[dict[str, Any]]:
    items: list[dict[str, Any]] = []
    if result is None:
        return items
    if isinstance(result, (list, tuple)):
        if len(result) >= 2 and isinstance(result[1], (list, tuple)) and result[1]:
            rect = _ocr_box_to_rect(result[0])
            text = str(result[1][0] or "").strip()
            if rect and text:
                return [
                    {
                        "page": page_number,
                        "x0": rect[0],
                        "y0": rect[1],
                        "x1": rect[2],
                        "y1": rect[3],
                        "text": text,
                        "page_width": width,
                        "page_height": height,
                    }
                ]
        for entry in result:
            items.extend(_extract_ocr_items_from_result(entry, page_number, width, height))
        return items

    mapping = result if isinstance(result, dict) else result if hasattr(result, "get") else None
    if mapping is not None:
        try:
            texts = mapping.get("rec_texts")
            if texts is None:
                texts = mapping.get("texts")
            boxes = mapping.get("rec_boxes")
            if boxes is None:
                boxes = mapping.get("rec_polys")
            if boxes is None:
                boxes = mapping.get("dt_polys")
            texts = texts or []
            boxes = boxes if boxes is not None else []
        except Exception:
            texts, boxes = [], []
        if len(texts) and len(boxes):
            for text, box in zip(texts, boxes):
                rect = _ocr_box_to_rect(box)
                cleaned = str(text or "").replace("\xa0", " ").strip()
                if cleaned and rect is not None:
                    items.append(
                        {
                            "page": page_number,
                            "x0": rect[0],
                            "y0": rect[1],
                            "x1": rect[2],
                            "y1": rect[3],
                            "text": cleaned,
                            "page_width": width,
                            "page_height": height,
                        }
                    )
            return items
    return items


def _ocr_box_to_rect(box: Any) -> tuple[float, float, float, float] | None:
    try:
        values = box.tolist() if hasattr(box, "tolist") else box
        if not values:
            return None
        if len(values) == 4 and all(isinstance(value, (int, float)) for value in values):
            x0, y0, x1, y1 = [float(value) for value in values]
            return (min(x0, x1), min(y0, y1), max(x0, x1), max(y0, y1))
        points = []
        for point in values:
            if isinstance(point, (list, tuple)) and len(point) >= 2:
                points.append((float(point[0]), float(point[1])))
        if not points:
            return None
        xs = [point[0] for point in points]
        ys = [point[1] for point in points]
        return (min(xs), min(ys), max(xs), max(ys))
    except Exception:
        return None


def _parse_chinese_pdf_layout(source: dict[str, Any], file_name: str) -> list[dict[str, Any]]:
    """Parse text-based Chinese timetable PDFs by word coordinates.

    The recovered text order from these PDF tables is often row-major but not
    cell-safe.  Using word coordinates lets us map each token to its real
    weekday/period column before building occupancy blocks.
    """

    doc = _open_pdf_document(source)
    if doc is None:
        return []

    name = _extract_chinese_name(_extract_pdf_text(source), file_name)
    all_blocks: list[dict[str, Any]] = []
    try:
        for page in doc:
            all_blocks.extend(_parse_chinese_table_page_items(_page_words(page), name, float(page.rect.height)))
    finally:
        doc.close()

    return all_blocks


def _parse_chinese_ocr_table_layout(source: dict[str, Any], file_name: str) -> list[dict[str, Any]]:
    items = _extract_pdf_ocr_items(source)
    if not items:
        return []
    name = _extract_chinese_name("", file_name)
    blocks: list[dict[str, Any]] = []
    for page in sorted({int(item.get("page", 0)) for item in items}):
        page_items = [
            {
                "x0": float(item.get("x0", 0)),
                "y0": float(item.get("y0", 0)),
                "x1": float(item.get("x1", 0)),
                "y1": float(item.get("y1", 0)),
                "text": _normalize_ocr_text(str(item.get("text", ""))),
            }
            for item in items
            if int(item.get("page", 0)) == page and str(item.get("text", "")).strip()
        ]
        page_height = max((float(item.get("page_height", 0)) for item in items if int(item.get("page", 0)) == page), default=0)
        if _detect_schedule_layout_profile_from_text_items("中方", "", page_items) == "cdut_undergrad_full_term_cn":
            specialized = _parse_chinese_full_term_grid_page_items(page_items, name, page_height)
            if specialized:
                blocks.extend(specialized)
                continue
        blocks.extend(_parse_chinese_table_page_items(page_items, name, page_height))
    return blocks


def _parse_chinese_table_page_items(
    items: list[dict[str, Any]],
    name: str,
    page_height: float,
) -> list[dict[str, Any]]:
    if not items:
        return []

    slot_labels = ["1-2", "3-4", "午", "5-6", "7-8", "9-11"]
    header_candidates = [
        item
        for item in items
        if item["text"] in slot_labels and _item_center(item)[1] < max(170.0, page_height * 0.28 if page_height else 170.0)
    ]
    header_groups = _cluster_items_by_y(header_candidates, 5.0)
    header_words = max(header_groups, key=len, default=[])
    if len(header_words) < len(slot_labels) * 5:
        return []

    header_words = sorted(header_words, key=lambda item: _item_center(item)[0])
    day_count = max(5, min(7, len(header_words) // len(slot_labels)))
    header_words = header_words[: day_count * len(slot_labels)]
    expected_slots = [(weekday, label) for weekday in WEEKDAYS[:day_count] for label in slot_labels]
    columns = [
        {"center": _item_center(item)[0], "weekday": weekday, "label": label}
        for item, (weekday, label) in zip(header_words, expected_slots)
    ]
    bands = _x_bands(columns)
    if not bands:
        return []
    header_bottom = max(float(item["y1"]) for item in header_words)

    week_candidates = []
    for item in items:
        match = re.fullmatch(r"(?:第)?(\d{1,2})周", item["text"])
        if match and _item_center(item)[0] < bands[0]["left"]:
            week_candidates.append({**item, "week": int(match.group(1))})
    week_rows = []
    for group in _cluster_items_by_y(week_candidates, 6.0):
        chosen = min(group, key=lambda item: item["x0"])
        week_rows.append(chosen)
    week_rows = sorted(week_rows, key=lambda item: _item_center(item)[1])
    if not week_rows:
        return []

    centers_y = [_item_center(item)[1] for item in week_rows]
    gaps = [b - a for a, b in zip(centers_y, centers_y[1:]) if b > a]
    default_gap = max(18.0, min(gaps) if gaps else 26.0)
    blocks: list[dict[str, Any]] = []
    for index, week_item in enumerate(week_rows):
        row_center = centers_y[index]
        row_top = max(header_bottom, (centers_y[index - 1] + row_center) / 2 if index else row_center - default_gap / 2)
        row_bottom = (
            (row_center + centers_y[index + 1]) / 2
            if index + 1 < len(centers_y)
            else min(page_height or row_center + default_gap, row_center + default_gap / 2)
        )
        if row_bottom <= row_top:
            row_bottom = row_top + default_gap

        row_words = [
            item
            for item in items
            if row_top <= _item_center(item)[1] < row_bottom
        ]
        date_text = " ".join(
            item["text"]
            for item in sorted(row_words, key=lambda value: (value["y0"], value["x0"]))
            if _item_center(item)[0] < bands[0]["left"] and re.search(r"\d{2}/\d{2}", item["text"])
        )
        week_start = _parse_week_start(date_text) if date_text else None
        cells: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
        for item in row_words:
            token = item["text"]
            if token in slot_labels or re.fullmatch(r"(?:第)?\d{1,2}周", token) or re.search(r"\d{2}/\d{2}", token):
                continue
            band = _pick_x_band(bands, _item_center(item)[0])
            if not band:
                continue
            cells[(band["weekday"], band["label"])].append(item)

        for (weekday, label), cell_words in cells.items():
            if label == "午":
                continue
            cell_text = _normalize_cell_text(cell_words)
            if not cell_text or _is_empty_or_non_class(cell_text):
                continue
            if _looks_like_room(cell_text):
                continue
            day = week_start + timedelta(days=WEEKDAY_INDEX.get(weekday, 0)) if week_start else _date_for_weekday(week_item["week"], weekday)
            for period in PERIOD_GROUPS.get(label, []):
                blocks.append(_block(name, "中方", week_item["week"], day, weekday, period, cell_text))
    return blocks


def _parse_chinese_full_term_grid_page_items(
    items: list[dict[str, Any]],
    name: str,
    page_height: float,
) -> list[dict[str, Any]]:
    if not items:
        return []

    week_rows = _find_chinese_full_term_week_rows(items)
    if len(week_rows) < 8:
        return []
    grid = _infer_chinese_full_term_grid(items, week_rows, page_height)
    if not grid:
        return []
    legend_aliases = _extract_chinese_full_term_legend_aliases(items, week_rows, grid)

    blocks: list[dict[str, Any]] = []
    for index, week_row in enumerate(week_rows):
        row_top = grid["header_bottom"] if index == 0 else (week_rows[index - 1]["y"] + week_row["y"]) / 2
        row_bottom = (week_row["y"] + week_rows[index + 1]["y"]) / 2 if index + 1 < len(week_rows) else week_row["y"] + grid["row_gap"] * 0.55
        if row_bottom <= row_top:
            row_bottom = week_row["y"] + grid["row_gap"] * 0.55

        row_items = [item for item in items if row_top <= _item_center(item)[1] < row_bottom]
        week_start = _parse_week_start(week_row.get("date_text", "")) if week_row.get("date_text") else None
        for cell in grid["cells"]:
            cell_items = [
                item
                for item in row_items
                if cell["left"] <= _item_center(item)[0] < cell["right"]
                and not _is_chinese_full_term_axis_token(str(item.get("text") or ""))
            ]
            if not cell_items:
                continue
            cell_text = _normalize_cell_text(cell_items)
            if not _is_chinese_full_term_course_text(cell_text, legend_aliases):
                continue
            day = (
                week_start + timedelta(days=WEEKDAY_INDEX.get(cell["weekday"], 0))
                if week_start
                else _date_for_weekday(week_row["week"], cell["weekday"])
            )
            for period in PERIOD_GROUPS.get(cell["label"], []):
                blocks.append(_block(name, "中方", week_row["week"], day, cell["weekday"], period, cell_text))
    return blocks


def _parse_chinese_full_term_visual_occupancy(source: dict[str, Any], file_name: str) -> list[dict[str, Any]]:
    doc = _open_pdf_document(source)
    if doc is None:
        return []
    try:
        import cv2
        import numpy as np
    except Exception:
        doc.close()
        return []

    try:
        name = _extract_name_from_filename(file_name) or Path(file_name).stem
        ocr_items_by_page: dict[int, list[dict[str, Any]]] = defaultdict(list)
        try:
            for item in _extract_pdf_ocr_items(source):
                ocr_items_by_page[int(item.get("page", 0))].append(item)
        except Exception:
            ocr_items_by_page = defaultdict(list)

        blocks: list[dict[str, Any]] = []
        for page_number, page in enumerate(doc):
            pix = page.get_pixmap(dpi=180, alpha=False)
            image = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)
            if pix.n > 3:
                image = image[:, :, :3]
            gray = cv2.cvtColor(image, cv2.COLOR_RGB2GRAY) if image.ndim == 3 else image
            blocks.extend(_parse_chinese_full_term_visual_image(gray, name, ocr_items_by_page.get(page_number, [])))
        return blocks
    finally:
        doc.close()


def _parse_chinese_full_term_visual_image(gray: Any, name: str, ocr_items: list[dict[str, Any]] | None = None) -> list[dict[str, Any]]:
    try:
        import cv2
        import numpy as np
    except Exception:
        return []
    if gray is None:
        return []

    inverted = cv2.threshold(gray, 200, 255, cv2.THRESH_BINARY_INV)[1]
    height, width = inverted.shape[:2]
    horizontal = cv2.morphologyEx(
        inverted,
        cv2.MORPH_OPEN,
        cv2.getStructuringElement(cv2.MORPH_RECT, (max(30, width // 80), 1)),
    )
    vertical = cv2.morphologyEx(
        inverted,
        cv2.MORPH_OPEN,
        cv2.getStructuringElement(cv2.MORPH_RECT, (1, max(20, height // 80))),
    )
    text_mask = cv2.subtract(inverted, cv2.bitwise_or(horizontal, vertical))

    row_lines = _cluster_int_positions(np.where((horizontal.sum(axis=1) // 255) > width * 0.18)[0], gap=5)
    col_lines = _cluster_int_positions(np.where((vertical.sum(axis=0) // 255) > height * 0.08)[0], gap=5)
    if len(row_lines) < 24 or len(col_lines) < 30:
        return []

    row_centers = [line[2] for line in row_lines]
    col_centers = [line[2] for line in col_lines]
    week_bounds = [row_centers[2], *row_centers[3:24]]
    if len(week_bounds) < 22:
        return []

    day_left = col_centers[1]
    day_right = col_centers[-1]
    if day_right <= day_left:
        return []

    ocr_items = ocr_items or []
    x_scale, y_scale = _ocr_to_visual_scale(ocr_items, width, height)
    blocks: list[dict[str, Any]] = []
    slot_units = [("1-2", 2), ("3-4", 2), ("午", 1), ("5-6", 2), ("7-8", 2), ("9-11", 3)]
    unit_total = sum(unit for _label, unit in slot_units)
    day_width = (day_right - day_left) / 7

    for week in range(1, 22):
        row_top = week_bounds[week - 1] + 2
        row_bottom = week_bounds[week] - 2
        if row_bottom <= row_top:
            continue
        for day_index, weekday in enumerate(WEEKDAYS):
            day_start = day_left + day_index * day_width
            day_end = day_left + (day_index + 1) * day_width
            cursor = day_start
            unit_width = (day_end - day_start) / unit_total
            for label, units in slot_units:
                next_cursor = cursor + unit_width * units
                if label != "午" and _visual_cell_has_occupancy(
                    text_mask,
                    int(cursor),
                    int(row_top),
                    int(next_cursor),
                    int(row_bottom),
                    ocr_items,
                    x_scale,
                    y_scale,
                ):
                    day = _date_for_weekday(week, weekday)
                    for period in PERIOD_GROUPS.get(label, []):
                        blocks.append(_block(name, "中方", week, day, weekday, period, "视觉表格占用"))
                cursor = next_cursor
    return blocks


def _cluster_int_positions(values: Any, gap: int = 5) -> list[tuple[int, int, int]]:
    positions = [int(value) for value in values]
    if not positions:
        return []
    clusters: list[tuple[int, int, int]] = []
    start = previous = positions[0]
    for value in positions[1:]:
        if value - previous <= gap:
            previous = value
        else:
            clusters.append((start, previous, (start + previous) // 2))
            start = previous = value
    clusters.append((start, previous, (start + previous) // 2))
    return clusters


def _ocr_to_visual_scale(ocr_items: list[dict[str, Any]], width: int, height: int) -> tuple[float, float]:
    page_width = max((float(item.get("page_width", 0)) for item in ocr_items), default=0.0)
    page_height = max((float(item.get("page_height", 0)) for item in ocr_items), default=0.0)
    return (
        width / page_width if page_width else 1.0,
        height / page_height if page_height else 1.0,
    )


def _visual_cell_has_occupancy(
    text_mask: Any,
    left: int,
    top: int,
    right: int,
    bottom: int,
    ocr_items: list[dict[str, Any]],
    x_scale: float,
    y_scale: float,
) -> bool:
    if right - left < 6 or bottom - top < 6:
        return False
    cell_texts = [
        _normalize_ocr_text(str(item.get("text") or ""))
        for item in ocr_items
        if left <= _item_center(item)[0] * x_scale < right and top <= _item_center(item)[1] * y_scale < bottom
    ]
    if cell_texts and all(_is_empty_or_non_class(text) or _is_chinese_full_term_axis_token(text) for text in cell_texts):
        return False

    crop = text_mask[top + 2 : bottom - 2, left + 3 : right - 3]
    if crop.size <= 0:
        return False
    dark_pixels = int((crop > 0).sum())
    ratio = dark_pixels / max(1, crop.size)
    return dark_pixels >= 35 and ratio >= 0.035


def _find_chinese_full_term_week_rows(items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    week_items = []
    for item in items:
        text = str(item.get("text") or "")
        match = re.search(r"(?<!\d)(\d{1,2})\s*[周期同]", text)
        if match and _item_center(item)[0] < 105:
            week = int(match.group(1))
            if 1 <= week <= 21:
                week_items.append({**item, "week": week})
    for group in _cluster_items_by_y(week_items, 7.0):
        chosen = min(group, key=lambda item: item["x0"])
        row_words = [
            item
            for item in items
            if abs(_item_center(item)[1] - _item_center(chosen)[1]) <= 12 and _item_center(item)[0] < 115
        ]
        date_text = " ".join(str(item.get("text") or "") for item in sorted(row_words, key=lambda value: value["x0"]))
        rows.append({"week": int(chosen["week"]), "y": _item_center(chosen)[1], "date_text": date_text})
    dedup: dict[int, dict[str, Any]] = {}
    for row in rows:
        dedup.setdefault(int(row["week"]), row)
    return sorted(dedup.values(), key=lambda row: row["y"])


def _infer_chinese_full_term_grid(
    items: list[dict[str, Any]],
    week_rows: list[dict[str, Any]],
    page_height: float,
) -> dict[str, Any] | None:
    header_items = [
        item
        for item in items
        if 85 <= _item_center(item)[1] <= 135 and _item_center(item)[0] > 85
    ]
    if not header_items:
        return None
    left = min(float(item["x0"]) for item in header_items)
    right_candidates = [
        float(item["x1"])
        for item in items
        if 90 <= _item_center(item)[0] and _item_center(item)[1] < min(page_height or 10_000, 570)
    ]
    right = max(right_candidates, default=0.0)
    if right <= left + 420:
        return None

    week_ys = [float(row["y"]) for row in week_rows]
    gaps = [b - a for a, b in zip(week_ys, week_ys[1:]) if b > a]
    row_gap = max(16.0, min(gaps) if gaps else 22.0)
    header_bottom = min(week_ys) - row_gap * 0.45
    day_width = (right - left) / 7
    units = [("1-2", 2), ("3-4", 2), ("午", 1), ("5-6", 2), ("7-8", 2), ("9-11", 3)]
    cells: list[dict[str, Any]] = []
    for day_index, weekday in enumerate(WEEKDAYS):
        day_left = left + day_index * day_width
        day_right = left + (day_index + 1) * day_width
        cursor = day_left
        unit_width = (day_right - day_left) / sum(width for _label, width in units)
        for label, width in units:
            next_cursor = cursor + unit_width * width
            cells.append({"weekday": weekday, "label": label, "left": cursor, "right": next_cursor})
            cursor = next_cursor
    return {"left": left, "right": right, "header_bottom": header_bottom, "row_gap": row_gap, "cells": cells}


def _extract_chinese_full_term_legend_aliases(
    items: list[dict[str, Any]],
    week_rows: list[dict[str, Any]],
    grid: dict[str, Any],
) -> set[str]:
    if not week_rows:
        return set()
    legend_top = max(float(row["y"]) for row in week_rows) + float(grid.get("row_gap", 20.0)) * 0.65
    legend_tokens = [
        str(item.get("text") or "").strip()
        for item in sorted(items, key=lambda value: (float(value.get("y0", 0)), float(value.get("x0", 0))))
        if _item_center(item)[1] >= legend_top and str(item.get("text") or "").strip()
    ]
    joined = " ".join(legend_tokens)
    aliases: set[str] = set()
    for match in re.finditer(r"[（(]\s*([A-Za-z0-9\u4e00-\u9fff]{1,4})\s*[)）]", joined):
        alias = match.group(1).strip()
        if _is_plausible_course_alias(alias):
            aliases.add(alias)

    # Some OCR chunks lose the surrounding parentheses but preserve common course names.
    known_course_aliases = {
        "高等数学": "高",
        "中国近现代史纲要": "中",
        "大学英语": "大",
        "大学体育": "大",
        "大学生心理健康": "大",
        "国家安全教育": "国",
        "形势与政策": "形",
        "思想政治": "思",
        "写作与演讲": "写",
        "人工智能": "人",
        "Python": "P",
        "社会重义发展史": "社",
        "社会主义发展史": "社",
    }
    for name, alias in known_course_aliases.items():
        if name in joined:
            aliases.add(alias)
    return aliases


def _is_plausible_course_alias(alias: str) -> bool:
    if not alias:
        return False
    if alias in {"校区", "主校区", "备注", "理", "实", "师", "室", "时", "学分"}:
        return False
    return bool(re.search(r"[\u4e00-\u9fffA-Za-z]", alias))


def _is_chinese_full_term_axis_token(token: str) -> bool:
    text = str(token or "").strip()
    if not text:
        return True
    if re.search(r"(?:学号|姓名|班级|学院|专业|生成日期|年级|周[/节]?|星期|成都理工大学)", text):
        return True
    if re.search(r"(?<!\d)\d{1,2}\s*[周期同]", text):
        return True
    if re.search(r"\d{2}/\d{2}", text):
        return True
    if _looks_like_chinese_slot_header_token(text):
        return True
    return False


def _looks_like_chinese_slot_header_token(token: str) -> bool:
    compact = re.sub(r"\s+", "", str(token or ""))
    if not compact:
        return False
    normalized = compact.replace("年", "午").replace("午", "午").replace("·", "-").replace("+", "-")
    return bool(
        re.search(r"1-?2", normalized)
        and (re.search(r"3-?4", normalized) or "34" in normalized)
        and ("午" in normalized or "5-6" in normalized or "56" in normalized)
    )


def _is_chinese_full_term_course_text(token: str, legend_aliases: set[str] | None = None) -> bool:
    text = str(token or "").strip()
    compact = re.sub(r"\s+", "", text)
    if _is_empty_or_non_class(compact) or _looks_like_room(compact):
        return False
    if re.search(r"(报到注册|运动会|运动食|清明节|请明节|劳动节|劳功节|劳助节|端午节|烯午节|考试周|全国英语四六级|不排|教研|教妍|微研|教好)", compact):
        return False
    if len(compact) <= 1:
        return False
    if re.fullmatch(r"[A-Za-z0-9:：/._-]+", compact) and not re.search(r"[A-Za-z]\d|\d[A-Za-z]", compact):
        return False
    aliases = legend_aliases or set()
    if aliases:
        if not any(_alias_matches_cell_text(alias, compact) for alias in aliases):
            return False
    return bool(re.search(r"[\u4e00-\u9fffA-Za-z]", compact))


def _alias_matches_cell_text(alias: str, compact: str) -> bool:
    if not alias or not compact:
        return False
    if alias in compact:
        return True
    if len(alias) == 1 and re.search(rf"{re.escape(alias)}[A-Za-z]?\d", compact):
        return True
    if alias.upper() in compact.upper() and re.search(r"[A-Za-z]\d", compact):
        return True
    return False


def _normalize_cell_text(cell_words: list[dict[str, Any]]) -> str:
    words = sorted(cell_words, key=lambda item: (round(item["y0"], 1), item["x0"]))
    pieces = [word["text"].strip() for word in words if word["text"].strip()]
    if not pieces:
        return ""
    return " ".join(pieces).strip()


def _parse_chinese_legacy_web_layout(source: dict[str, Any], file_name: str) -> list[dict[str, Any]]:
    """Parse old jw.cdut.edu.cn timetable-print PDFs by coordinates."""

    path = _source_path(source)
    if not path or not Path(path).exists():
        return []

    try:
        import fitz
    except Exception:
        return []

    try:
        doc = fitz.open(path)
    except Exception:
        return []

    text = _extract_pdf_text(source)
    if "学期理论课表" not in text and "旧版课表打印" not in text:
        return []

    name = _extract_chinese_name(text, file_name)
    blocks: list[dict[str, Any]] = []
    for page in doc:
        words = [
            {
                "x0": float(word[0]),
                "y0": float(word[1]),
                "x1": float(word[2]),
                "y1": float(word[3]),
                "text": str(word[4]).replace("\xa0", " ").strip(),
            }
            for word in page.get_text("words")
            if str(word[4]).replace("\xa0", " ").strip()
        ]
        header_words = sorted(
            [word for word in words if word["text"] in ("星期一", "星期二", "星期三", "星期四", "星期五", "星期六", "星期日")],
            key=lambda word: word["x0"],
        )
        if len(header_words) < 5:
            continue

        columns = []
        for word in header_words:
            weekday = word["text"].replace("星期", "周")
            columns.append({"center": (word["x0"] + word["x1"]) / 2, "weekday": weekday})
        header_y = min(word["y0"] for word in header_words)

        day_words: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for word in words:
            if word["y0"] <= header_y + 8:
                continue
            if word["x0"] < min(column["center"] for column in columns) - 45:
                continue
            x_center = (word["x0"] + word["x1"]) / 2
            nearest = min(columns, key=lambda column: abs(column["center"] - x_center))
            if abs(nearest["center"] - x_center) > 50:
                continue
            token = word["text"]
            if not token or (token and 0xE000 <= ord(token[0]) <= 0xF8FF) or token in {"---------------------"}:
                continue
            day_words[nearest["weekday"]].append(word)

        for weekday, column_words in day_words.items():
            lines = [word["text"] for word in sorted(column_words, key=lambda item: (item["y0"], item["x0"]))]
            for course, week_expr, period_expr in _extract_legacy_course_specs(lines):
                weeks = _expand_week_expr(week_expr)
                periods = _expand_period_expr(period_expr)
                for week in weeks:
                    day = _date_for_weekday(week, weekday)
                    for period in periods:
                        if 1 <= period <= 11:
                            blocks.append(_block(name, "中方", week, day, weekday, period, course))

    return blocks


def _parse_chinese_ocr_legacy_layout(source: dict[str, Any], file_name: str) -> list[dict[str, Any]]:
    items = _extract_pdf_ocr_items(source)
    if not items:
        return []
    text = "\n".join(str(item.get("text", "")) for item in items)
    if not any(keyword in text for keyword in ("学期理论课表", "旧版课表打印", "星期一")):
        return []
    name = _extract_chinese_name(text, file_name)
    blocks: list[dict[str, Any]] = []
    for page in sorted({int(item.get("page", 0)) for item in items}):
        page_items = [
            {
                "x0": float(item.get("x0", 0)),
                "y0": float(item.get("y0", 0)),
                "x1": float(item.get("x1", 0)),
                "y1": float(item.get("y1", 0)),
                "text": _normalize_ocr_text(str(item.get("text", ""))),
            }
            for item in items
            if int(item.get("page", 0)) == page and str(item.get("text", "")).strip()
        ]
        blocks.extend(_parse_chinese_legacy_page_items(page_items, name))
    return blocks


def _parse_chinese_legacy_page_items(items: list[dict[str, Any]], name: str) -> list[dict[str, Any]]:
    header_texts = {"星期一", "星期二", "星期三", "星期四", "星期五", "星期六", "星期日"}
    header_words = sorted(
        [item for item in items if item["text"] in header_texts],
        key=lambda item: item["x0"],
    )
    if len(header_words) < 5:
        return []

    columns = []
    for item in header_words:
        columns.append({"center": (item["x0"] + item["x1"]) / 2, "weekday": item["text"].replace("星期", "周")})
    centers = sorted(column["center"] for column in columns)
    min_gap = min((b - a for a, b in zip(centers, centers[1:])), default=90)
    max_distance = max(45.0, min_gap * 0.52)
    header_y = min(item["y0"] for item in header_words)

    day_words: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for item in items:
        if item["y0"] <= header_y + 10:
            continue
        token = item["text"]
        if not token or token in header_texts or token == "---------------------":
            continue
        if token and 0xE000 <= ord(token[0]) <= 0xF8FF:
            continue
        x_center = (item["x0"] + item["x1"]) / 2
        nearest = min(columns, key=lambda column: abs(column["center"] - x_center))
        if abs(nearest["center"] - x_center) > max_distance:
            continue
        day_words[nearest["weekday"]].append(item)

    blocks: list[dict[str, Any]] = []
    for weekday, column_words in day_words.items():
        lines = [item["text"] for item in sorted(column_words, key=lambda value: (value["y0"], value["x0"]))]
        for course, week_expr, period_expr in _extract_legacy_course_specs(lines):
            for week in _expand_week_expr(week_expr):
                day = _date_for_weekday(week, weekday)
                for period in _expand_period_expr(period_expr):
                    if 1 <= period <= 11:
                        blocks.append(_block(name, "中方", week, day, weekday, period, course))
    return blocks


def _normalize_ocr_text(text: str) -> str:
    return re.sub(r"\s+", "", text.replace("\xa0", " ")).strip()


def _extract_legacy_course_specs(lines: list[str]) -> list[tuple[str, str, str]]:
    specs: list[tuple[str, str, str]] = []
    for index, line in enumerate(lines):
        combined = line
        if "(周)" not in combined and index + 1 < len(lines):
            combined = f"{combined}{lines[index + 1]}"
        if "[控制]" in combined:
            continue
        if "(周)" in combined and "[" not in combined and index + 1 < len(lines):
            combined = f"{combined}{lines[index + 1]}"
        if "[0" in combined and "节]" not in combined and index + 1 < len(lines):
            combined = f"{combined}{lines[index + 1]}"
        match = re.search(r"([0-9,，、\.\-\s]+)\(周\)\s*\[([0-9\-\s]+)节\]", combined)
        if not match:
            continue
        if index > 0 and re.fullmatch(r"[0-9\-\s]+节\]?", line):
            continue
        course_parts: list[str] = []
        cursor = index - 2
        while cursor >= 0 and len(course_parts) < 3:
            token = lines[cursor].strip()
            if not token or token == "---------------------":
                break
            if re.search(r"\(周\)|\[.*节\]|^[A-Z]\d|^E\d|^C\d|^校外", token):
                break
            if not re.fullmatch(r"[\u4e00-\u9fff]{2,4}", token):
                course_parts.insert(0, token)
            cursor -= 1
        course = "".join(course_parts).strip() or "课程"
        specs.append((course, match.group(1), match.group(2)))
    return specs


def _expand_week_expr(expr: str) -> list[int]:
    weeks: set[int] = set()
    for part in re.split(r"[,，、\.\s]+", expr.strip()):
        if not part:
            continue
        if "-" in part:
            start, end = part.split("-", 1)
            if start.strip().isdigit() and end.strip().isdigit():
                weeks.update(range(int(start), int(end) + 1))
        elif part.isdigit():
            weeks.add(int(part))
    return sorted(week for week in weeks if 1 <= week <= 30)


def _expand_period_expr(expr: str) -> list[int]:
    numbers = [int(value) for value in re.findall(r"\d{1,2}", expr)]
    return sorted({number for number in numbers if 1 <= number <= 12})


def _parse_chinese_text(text: str, file_name: str) -> list[dict[str, Any]]:
    name = _extract_chinese_name(text, file_name)
    lines = _clean_lines(text)
    days = [weekday for weekday in WEEKDAYS if weekday.replace("周", "星期") in lines or weekday in lines]
    if not days:
        days = WEEKDAYS[:5]
    slots = [(weekday, group) for weekday in days for group in ["1-2", "3-4", "午", "5-6", "7-8", "9-11"]]

    blocks: list[dict[str, Any]] = []
    index = 0
    while index < len(lines):
        week_match = re.match(r"^(\d{1,2})周$", lines[index])
        if not week_match:
            index += 1
            continue
        week = int(week_match.group(1))
        week_start = None
        if index + 1 < len(lines):
            week_start = _parse_week_start(lines[index + 1])
        index += 2
        slot_index = 0
        while index < len(lines) and not re.match(r"^\d{1,2}周$", lines[index]):
            token = lines[index]
            if slot_index >= len(slots):
                index += 1
                continue
            weekday, group = slots[slot_index]
            if _is_empty_or_non_class(token):
                slot_index += 1
                index += 1
                continue
            if _looks_like_room(token):
                index += 1
                continue
            course = token
            if index + 1 < len(lines) and _looks_like_room(lines[index + 1]):
                index += 1
            if group != "午":
                for period in PERIOD_GROUPS[group]:
                    day = week_start + timedelta(days=WEEKDAY_INDEX.get(weekday, 0)) if week_start else _date_for_weekday(week, weekday)
                    blocks.append(_block(name, "中方", week, day, weekday, period, course))
            slot_index += 1
            index += 1
    return blocks


def _parse_english_ocr_week_grid(source: dict[str, Any], file_name: str) -> list[dict[str, Any]]:
    items = _extract_pdf_ocr_items(source)
    if not items:
        return []
    name = _extract_chinese_name("", file_name)
    date_grid_blocks: list[dict[str, Any]] = []
    fallback_columns: list[dict[str, Any]] | None = None
    for page in sorted({int(item.get("page", 0)) for item in items}):
        page_items = [
            {
                "x0": float(item.get("x0", 0)),
                "y0": float(item.get("y0", 0)),
                "x1": float(item.get("x1", 0)),
                "y1": float(item.get("y1", 0)),
                "text": _normalize_ocr_text(str(item.get("text", ""))),
            }
            for item in items
            if int(item.get("page", 0)) == page and str(item.get("text", "")).strip()
        ]
        headers = _find_english_period_header_rows(page_items)
        if headers:
            fallback_columns = headers[-1]["columns"]
        elif fallback_columns and _find_english_day_rows(page_items):
            page_items.extend(
                {
                    "x0": float(column["center"]) - 2,
                    "y0": 0.0,
                    "x1": float(column["center"]) + 2,
                    "y1": 8.0,
                    "text": str(column["period"]),
                }
                for column in fallback_columns
            )
        date_grid_blocks.extend(_parse_english_grid_page_items_from_pdf_words(page_items, name))
    if date_grid_blocks:
        return date_grid_blocks

    normalized_text = "\n".join(_normalize_ocr_text(str(item.get("text", ""))) for item in items)
    if "第" not in normalized_text or "周" not in normalized_text or "周一" not in normalized_text:
        return []

    blocks: list[dict[str, Any]] = []
    for page in sorted({int(item.get("page", 0)) for item in items}):
        page_items = [
            {
                "x0": float(item.get("x0", 0)),
                "y0": float(item.get("y0", 0)),
                "x1": float(item.get("x1", 0)),
                "y1": float(item.get("y1", 0)),
                "text": _normalize_ocr_text(str(item.get("text", ""))),
            }
            for item in items
            if int(item.get("page", 0)) == page and str(item.get("text", "")).strip()
        ]
        blocks.extend(_parse_english_week_grid_page_items(page_items, name))
    return blocks


def _parse_english_pdf_grid_layout(source: dict[str, Any], file_name: str) -> list[dict[str, Any]]:
    doc = _open_pdf_document(source)
    if doc is None:
        return []

    name = _extract_chinese_name("", file_name)
    blocks: list[dict[str, Any]] = []
    try:
        for page in doc:
            items = _page_words(page)
            if not items:
                continue
            blocks.extend(_parse_english_grid_page_items_from_pdf_words(items, name))
    finally:
        doc.close()
    return blocks


def _parse_english_grid_page_items_from_pdf_words(items: list[dict[str, Any]], name: str) -> list[dict[str, Any]]:
    header_rows = _find_english_period_header_rows(items)
    if not header_rows:
        return []
    day_rows = _find_english_day_rows(items)
    if not day_rows:
        return []

    blocks: list[dict[str, Any]] = []
    for index, day_row in enumerate(day_rows):
        row_y = day_row["y"]
        header = max((row for row in header_rows if row["y"] < row_y), key=lambda row: row["y"], default=None)
        if not header:
            continue
        previous_day_y = day_rows[index - 1]["y"] if index else None
        next_day_y = day_rows[index + 1]["y"] if index + 1 < len(day_rows) else None
        next_header_y = min((row["y"] for row in header_rows if row["y"] > row_y), default=row_y + 80)
        row_top = (previous_day_y + row_y) / 2 if previous_day_y is not None else row_y - 8
        row_bottom = (row_y + next_day_y) / 2 if next_day_y is not None else min(next_header_y - 2, row_y + 48)
        if row_bottom <= row_top:
            row_bottom = row_y + 30

        bands = _x_bands(header["columns"])
        if not bands:
            continue
        cells: dict[int, list[dict[str, Any]]] = defaultdict(list)
        for item in items:
            y_center = (item["y0"] + item["y1"]) / 2
            if not (row_top <= y_center <= row_bottom):
                continue
            x_center = (item["x0"] + item["x1"]) / 2
            token = item["text"]
            if _is_english_grid_noise_token(token):
                continue
            nearest = _pick_x_band(bands, x_center)
            if not nearest:
                continue
            cells[int(nearest["period"])].append(item)

        for period, cell_items in cells.items():
            tokens = [
                item["text"]
                for item in sorted(cell_items, key=lambda value: (round(value["y0"], 1), value["x0"]))
            ]
            course = _merge_english_grid_course_tokens(tokens)
            if not course:
                continue
            blocks.append(
                _block(
                    name,
                    "英方",
                    _week_from_date(day_row["date"]),
                    day_row["date"],
                    WEEKDAYS[day_row["date"].weekday()],
                    period,
                    course,
                )
            )
    return blocks


def _find_english_period_header_rows(items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    number_items: list[dict[str, Any]] = []
    for item in items:
        periods = _extract_english_period_header_numbers(str(item.get("text") or ""))
        if not periods:
            continue
        width = max(1.0, float(item["x1"]) - float(item["x0"]))
        for offset, period in enumerate(periods):
            left = float(item["x0"]) + width * offset / len(periods)
            right = float(item["x0"]) + width * (offset + 1) / len(periods)
            number_items.append({**item, "text": str(period), "x0": left, "x1": right})
    grouped: list[list[dict[str, Any]]] = []
    for item in sorted(number_items, key=lambda value: value["y0"]):
        for group in grouped:
            group_y = sum(_item_center(value)[1] for value in group) / len(group)
            if abs(group_y - _item_center(item)[1]) <= 12.0:
                group.append(item)
                break
        else:
            grouped.append([item])

    rows: list[dict[str, Any]] = []
    for group in grouped:
        by_period: dict[int, dict[str, Any]] = {}
        for item in group:
            period = int(item["text"])
            by_period[period] = item
        columns = _complete_english_period_columns(by_period)
        if len(columns) < 8 or not {1, 2, 3, 4}.issubset({int(column["period"]) for column in columns}):
            continue
        rows.append(
            {
                "y": sum(_item_center(item)[1] for item in group) / len(group),
                "columns": columns,
                "direct_count": sum(1 for column in columns if not column.get("synthetic")),
            }
        )
    return _align_english_header_rows(sorted(rows, key=lambda row: row["y"]))


def _complete_english_period_columns(by_period: dict[int, dict[str, Any]]) -> list[dict[str, Any]]:
    known = {
        int(period): _item_center(item)[0]
        for period, item in by_period.items()
        if 1 <= int(period) <= 11
    }
    if len(known) >= 8:
        return [{"period": period, "center": known[period], "synthetic": False} for period in sorted(known)]
    if len(known) < 6 or not {1, 2, 3, 4}.issubset(known):
        return [{"period": period, "center": known[period], "synthetic": False} for period in sorted(known)]

    completed = dict(known)
    for period in range(1, 12):
        if period in completed:
            continue
        lower = max((value for value in known if value < period), default=None)
        upper = min((value for value in known if value > period), default=None)
        if lower is not None and upper is not None:
            span = upper - lower
            completed[period] = known[lower] + (known[upper] - known[lower]) * ((period - lower) / span)

    return [
        {"period": period, "center": completed[period], "synthetic": period not in known}
        for period in sorted(completed)
    ]


def _align_english_header_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    if len(rows) < 2:
        return rows
    templates = [
        row
        for row in rows
        if sum(1 for column in row["columns"] if not column.get("synthetic")) >= 8
    ]
    if not templates:
        return rows
    template = max(templates, key=lambda row: row.get("direct_count", 0))
    template_centers = {int(column["period"]): float(column["center"]) for column in template["columns"]}
    aligned: list[dict[str, Any]] = []
    for row in rows:
        columns = []
        for column in row["columns"]:
            period = int(column["period"])
            if column.get("synthetic") and period in template_centers:
                columns.append({**column, "center": template_centers[period]})
            else:
                columns.append(column)
        aligned.append({**row, "columns": columns})
    return aligned


def _extract_english_period_header_numbers(token: str) -> list[int]:
    text = str(token or "").strip()
    if not text:
        return []
    compact = re.sub(r"\s+", "", text)
    time_periods = _extract_english_period_time_header_numbers(compact)
    if time_periods:
        return time_periods
    if compact == "A":
        return [1]
    if re.fullmatch(r"[.。]*([1-9]|10|11)[.。]*", compact):
        return [int(re.sub(r"\D", "", compact))]
    numbers = [int(value) for value in re.findall(r"(?<!\d)(?:[1-9]|10|11)(?!\d)", compact)]
    if 2 <= len(numbers) <= 4:
        ordered = []
        for number in numbers:
            if 1 <= number <= 11 and number not in ordered:
                ordered.append(number)
        if len(ordered) < 2:
            return []
        if ordered == list(range(min(ordered), max(ordered) + 1)):
            return ordered
    return []


def _extract_english_period_time_header_numbers(token: str) -> list[int]:
    compact = re.sub(r"\s+", "", str(token or ""))
    if not compact or re.search(r"[A-Za-z\u4e00-\u9fff]", compact):
        return []
    digits = re.sub(r"\D", "", compact)
    if len(digits) < 3:
        return []

    if re.match(r"8[:：]?[01]\d", compact) or digits.startswith(("9810", "98105", "98195", "810")):
        return [1]
    if re.match(r"9[%:：]?[04]\d", compact) or digits.startswith(("9045", "09045", "085045", "0945", "045")):
        return [2]
    if re.match(r"10[:：]?\d", compact):
        return [3]
    if re.match(r"11[:：]?\d", compact):
        return [4]
    if re.match(r"14[:：]?\d", compact) or digits.startswith(("1430", "145", "153")):
        return [5]
    if re.match(r"15[:：]?20", compact) or digits.startswith(("1520", "15205")):
        return [6]
    if re.match(r"16[:：;]?\d", compact):
        return [7]
    if re.match(r"17[:：]?\d", compact) or digits.startswith(("178", "1715")):
        return [8]
    if re.match(r"19[%:：]?\d", compact) or digits.startswith(("1910", "1")) and "15" in digits:
        return [9]
    if re.match(r"20[:：]?[04]\d", compact) or digits.startswith(("2000", "20645", "2045")):
        return [10]
    if re.match(r"20[:：]?5\d", compact) or digits.startswith(("2050", "20155", "201505")):
        return [11]
    return []


def _find_english_day_rows(items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    weekday_items = [
        item
        for item in items
        if item["text"] in ENGLISH_WEEKDAY_MAP and item["x0"] < 95
    ]
    date_items = [
        item
        for item in items
        if _is_english_date_token(item["text"]) and item["x0"] < 105
    ]
    for weekday in weekday_items:
        candidates = [
            item
            for item in date_items
            if -4 <= item["y0"] - weekday["y0"] <= 18
        ]
        if not candidates:
            continue
        date_item = min(candidates, key=lambda item: abs(item["x0"] - weekday["x0"]))
        try:
            day = _parse_date(date_item["text"])
        except Exception:
            continue
        rows.append(
            {
                "weekday_text": weekday["text"],
                "date": day,
                "y": min(_item_center(weekday)[1], _item_center(date_item)[1]),
            }
        )
    date_pattern = r"(?:20\d{2}[-/]?\d{2}[-/]?\d{2}|20\d{4}-\d{2}|\d{2}/\d{2})"
    weekday_pattern = r"Monday|Tuesday|Wednesday|Thursday|Friday|Saturday|Sunday"
    for item in items:
        if item["x0"] >= 115:
            continue
        match = re.search(rf"({weekday_pattern}).*?({date_pattern})", item["text"])
        if not match:
            match = re.search(rf"({date_pattern}).*?({weekday_pattern})", item["text"])
        if not match:
            continue
        date_text = match.group(2) if match.group(1) in ENGLISH_WEEKDAY_MAP else match.group(1)
        weekday_text = match.group(1) if match.group(1) in ENGLISH_WEEKDAY_MAP else match.group(2)
        try:
            day = _parse_date(date_text)
        except Exception:
            continue
        rows.append({"weekday_text": weekday_text, "date": day, "y": _item_center(item)[1]})
    dedup: dict[date, dict[str, Any]] = {}
    for row in rows:
        dedup.setdefault(row["date"], row)
    return sorted(dedup.values(), key=lambda row: row["y"])


def _parse_english_week_grid_page_items(items: list[dict[str, Any]], name: str) -> list[dict[str, Any]]:
    week_titles = sorted(
        [item for item in items if re.search(r"第\s*(\d{1,2})\s*周", item["text"])],
        key=lambda item: (item["y0"], item["x0"]),
    )
    if not week_titles:
        return []

    blocks: list[dict[str, Any]] = []
    weekday_set = set(WEEKDAYS)
    time_periods = {
        "8:10": [1, 2],
        "10:15": [3, 4],
        "12:40": [12, 13],
        "13:30": [13],
        "14:30": [5, 6],
        "16:25": [7, 8],
        "19:10": [9, 10],
        "20:50": [11],
    }

    for title_index, title in enumerate(week_titles):
        match = re.search(r"第\s*(\d{1,2})\s*周", title["text"])
        if not match:
            continue
        week = int(match.group(1))
        x_min = title["x0"] - 100
        x_max = title["x0"] + 280
        next_same_column = [
            other
            for other in week_titles[title_index + 1 :]
            if other["y0"] > title["y0"] + 120 and abs(other["x0"] - title["x0"]) < 120
        ]
        y_top = title["y0"]
        y_bottom = min((other["y0"] - 18 for other in next_same_column), default=title["y0"] + 560)

        headers = sorted(
            [
                item
                for item in items
                if y_top <= item["y0"] <= y_top + 90
                and x_min <= (item["x0"] + item["x1"]) / 2 <= x_max
                and item["text"] in weekday_set
            ],
            key=lambda item: item["x0"],
        )
        if len(headers) < 5:
            continue

        columns = [{"center": (item["x0"] + item["x1"]) / 2, "weekday": item["text"]} for item in headers]
        centers = sorted(column["center"] for column in columns)
        min_gap = min((b - a for a, b in zip(centers, centers[1:])), default=34)
        column_tolerance = max(16.0, min_gap * 0.55)
        header_y = max(item["y1"] for item in headers)

        time_items = sorted(
            [
                item
                for item in items
                if header_y <= item["y0"] <= y_bottom
                and x_min <= item["x0"] <= x_min + 95
                and item["text"] in time_periods
            ],
            key=lambda item: item["y0"],
        )
        if not time_items:
            continue

        row_bounds: list[tuple[dict[str, Any], float, float]] = []
        for index, time_item in enumerate(time_items):
            center_y = (time_item["y0"] + time_item["y1"]) / 2
            previous_y = (time_items[index - 1]["y0"] + time_items[index - 1]["y1"]) / 2 if index else header_y
            next_y = (time_items[index + 1]["y0"] + time_items[index + 1]["y1"]) / 2 if index + 1 < len(time_items) else y_bottom
            row_bounds.append((time_item, (previous_y + center_y) / 2, (center_y + next_y) / 2))

        for time_item, row_top, row_bottom in row_bounds:
            row_periods = time_periods[time_item["text"]]
            cells: dict[str, list[str]] = defaultdict(list)
            for item in items:
                if not (row_top <= (item["y0"] + item["y1"]) / 2 < row_bottom):
                    continue
                x_center = (item["x0"] + item["x1"]) / 2
                if not (centers[0] - column_tolerance <= x_center <= centers[-1] + column_tolerance):
                    continue
                nearest = min(columns, key=lambda column: abs(column["center"] - x_center))
                if abs(nearest["center"] - x_center) > column_tolerance:
                    continue
                token = item["text"]
                if _is_english_grid_noise_token(token):
                    continue
                cells[nearest["weekday"]].append(token)

            for weekday, tokens in cells.items():
                course = _merge_english_grid_course_tokens(tokens)
                if not course:
                    continue
                day = _date_for_weekday(week, weekday)
                for period in row_periods:
                    blocks.append(_block(name, "英方", week, day, weekday, period, course))
    return blocks


def _is_english_grid_noise_token(token: str) -> bool:
    if not token or token in WEEKDAYS:
        return True
    if re.search(r"第\d{1,2}周", token):
        return True
    if _extract_english_period_header_numbers(token):
        return True
    if re.fullmatch(r"\d{1,4}", token) or re.fullmatch(r"\d{1,2}:\d{2}", token):
        return True
    if re.search(r"\d{1,2}[:.]\d{2}", token) and not re.search(r"[A-Za-z\u4e00-\u9fff]", token):
        return True
    if re.search(r"\d", token) and not re.search(r"[A-Za-z\u4e00-\u9fff]", token):
        return True
    if re.fullmatch(r"\d{1,2}月?", token):
        return True
    if re.match(r"^(?:\d?[A-Z]\d|E\d|TT\d|C\d|L\d|Room|Group|Teacher)", token, re.IGNORECASE):
        return True
    if token in {"Simpl", "Home", "Start", "End"}:
        return True
    return False


def _merge_english_grid_course_tokens(tokens: list[str]) -> str:
    clean = []
    for token in tokens:
        normalized = token.strip()
        if not normalized or _is_english_grid_noise_token(normalized):
            continue
        clean.append(normalized)
    if not clean:
        return ""
    course = "".join(clean) if any(re.search(r"[\u4e00-\u9fff]", token) for token in clean) else " ".join(clean)
    course = course.replace("形势与政策", "形势与政策").replace("改革开放史", "改革开放史")
    return course.strip()


def _parse_english_text(text: str, file_name: str) -> list[dict[str, Any]]:
    name = _extract_chinese_name("", file_name)
    lines = _clean_lines(text)
    blocks: list[dict[str, Any]] = []
    current_dates: list[tuple[str, date]] = []
    current_period = None
    i = 0
    while i < len(lines):
        line = lines[i]
        if line in ENGLISH_WEEKDAY_MAP and i + 1 < len(lines) and re.match(r"\d{4}-\d{2}-\d{2}", lines[i + 1]):
            current_dates.append((ENGLISH_WEEKDAY_MAP[line], _parse_date(lines[i + 1])))
            if len(current_dates) > 7:
                current_dates = current_dates[-7:]
            i += 2
            continue
        if re.fullmatch(r"(?:[1-9]|10|11)", line):
            current_period = int(line)
            i += 1
            continue
        if current_period and current_dates and _is_english_course_token(line):
            date_for_course = current_dates[-1][1]
            weekday = WEEKDAYS[date_for_course.weekday()]
            week = _week_from_date(date_for_course)
            blocks.append(_block(name, "英方", week, date_for_course, weekday, current_period, line))
            current_period = None
        i += 1
    return blocks


def _block(name: str, source: str, week: int, day: date, weekday: str, period: int, course: str) -> dict[str, Any]:
    return {
        "name": name,
        "source": source,
        "week": int(week),
        "date": day.isoformat(),
        "weekday": weekday,
        "period": int(period),
        "time": period_time_range(period),
        "course": course,
    }


def _clean_lines(text: str) -> list[str]:
    return [line.replace("\xa0", " ").strip() for line in text.splitlines()]


def _extract_name_from_filename(file_name: str) -> str:
    stem = Path(file_name).stem
    for part in re.split(r"[-_\s]+", stem):
        token = re.sub(r"[^\u4e00-\u9fff]", "", part)
        if re.fullmatch(r"[\u4e00-\u9fff]{2,4}", token) and token not in NAME_STOPWORDS:
            return token
    cleaned = stem
    for word in sorted(NAME_STOPWORDS, key=len, reverse=True):
        cleaned = cleaned.replace(word, " ")
    for token in re.findall(r"[\u4e00-\u9fff]{2,4}", cleaned):
        if token not in NAME_STOPWORDS:
            return token
    return ""


def _extract_chinese_name(text: str, file_name: str) -> str:
    filename_name = _extract_name_from_filename(file_name)
    if filename_name:
        return filename_name
    candidate = _extract_chinese_name_from_text(text)
    if candidate:
        return candidate
    stem = Path(file_name).stem
    return stem


def _extract_chinese_name_from_text(text: str) -> str:
    match = re.search(r"姓名\s*[:：]?\s*([\u4e00-\u9fff]{2,4})", text)
    if not match:
        return ""
    candidate = match.group(1)
    return "" if candidate in NAME_STOPWORDS else candidate


def _parse_week_start(text: str) -> date | None:
    match = re.search(r"(\d{2})/(\d{2})", text)
    if not match:
        return None
    return date(2026, int(match.group(1)), int(match.group(2)))


def _parse_date(text: str) -> date:
    text = text.strip()
    if re.match(r"\d{4}-\d{2}-\d{2}", text):
        return datetime.strptime(text[:10], "%Y-%m-%d").date()
    digits = re.sub(r"\D", "", text)
    if len(digits) >= 8 and digits.startswith("20"):
        return datetime.strptime(digits[:8], "%Y%m%d").date()
    if re.match(r"\d{2}/\d{2}", text):
        month, day = map(int, text[:5].split("/"))
        return date(2026, month, day)
    return date.fromisoformat(text[:10])


def _is_english_date_token(text: str) -> bool:
    try:
        day = _parse_date(str(text))
    except Exception:
        return False
    return 2020 <= day.year <= 2035


def _week_from_date(day: date) -> int:
    return max(1, ((day - _semester_start_date()).days // 7) + 1)


def _date_for_weekday(week: int, weekday: str) -> date:
    return _semester_start_date() + timedelta(days=(week - 1) * 7 + WEEKDAY_INDEX.get(weekday, 0))


def _school_calendar_config() -> dict[str, Any]:
    path = Path(__file__).resolve().parents[1] / "config" / "school_calendar.json"
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return {}


def _semester_start_date() -> date:
    configured = os.environ.get("KONGGU_SEMESTER_START_DATE")
    if configured:
        try:
            return date.fromisoformat(configured)
        except Exception:
            pass
    value = _school_calendar_config().get("semester_start_date")
    if value:
        try:
            return date.fromisoformat(str(value))
        except Exception:
            pass
    return SEMESTER_START


def _teaching_weeks() -> int:
    configured = os.environ.get("KONGGU_TEACHING_WEEKS")
    if configured:
        try:
            weeks = int(configured)
            if weeks > 0:
                return weeks
        except Exception:
            pass
    value = _school_calendar_config().get("teaching_weeks")
    try:
        weeks = int(value)
    except Exception:
        return TEACHING_WEEKS
    return weeks if weeks > 0 else TEACHING_WEEKS


def _is_empty_or_non_class(token: str) -> bool:
    text = str(token or "").strip()
    compact = re.sub(r"[\s:：;；,，。|｜/\\]+", "", text)
    if text in NON_CLASS_KEYWORDS or compact in {"", "午", "无", "空", "不排课", "暂无", "None", "N/A", "-", "校外", "教研"}:
        return True
    if any(
        keyword in compact
        for keyword in (
            "校外",
            "教研",
            "教妍",
            "微研",
            "教好",
            "不排课",
            "不排误",
            "不持课",
            "不排深",
            "清明节",
            "请明节",
            "劳动节",
            "劳功节",
            "劳助节",
            "端午节",
            "烯午节",
            "运动会",
            "运动食",
            "期末周",
            "考试周",
            "全国英语四六级",
        )
    ):
        return True
    if text.startswith("-"):
        return True
    if re.search(r"https?://|www\.", text, re.IGNORECASE):
        return True
    return False


def _looks_like_room(token: str) -> bool:
    compact = re.sub(r"\s+", "", token.strip())
    return bool(re.search(r"^(?:E\d|C\d|TT|内|大\d|L\d|[A-Z]\d|C075|Room|教室)", compact, re.IGNORECASE))


def _is_english_course_token(token: str) -> bool:
    if _is_empty_or_non_class(token) or re.match(r"\d{4}-\d{2}-\d{2}", token):
        return False
    if token in ENGLISH_WEEKDAY_MAP or re.fullmatch(r"(?:[1-9]|10|11)", token):
        return False
    if re.match(r"\d{2}:\d{2}", token) or token in {"Home", "Teacher:", "Room:", "Group:", "Start:", "End:", "---"}:
        return False
    return bool(re.search(r"[A-Za-z\u4e00-\u9fff]", token))
